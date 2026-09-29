from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from pydantic import TypeAdapter, ValidationError

from .models import (
    DatasetAsset,
    DatasetColumn,
    DiagramAsset,
    PictogramGridAsset,
    Scalar,
    StructuredAsset,
    VisualHint,
    VisualCapabilities,
    VisualSlotContent,
)


ASSET_LIST_ADAPTER = TypeAdapter(list[StructuredAsset])
SUPPORTED_VISUAL_SUBTYPES = {
    "chart": {"bar", "column", "line", "area", "pie", "donut", "scatter"},
    "diagram": {"process", "timeline", "cycle", "hierarchy", "comparison"},
}


class StructuredAssetError(ValueError):
    pass


def _slug(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9]+", "-", value.strip().lower()).strip("-")
    return value[:48] or "dataset"


def _stable_id(title: str, columns: Iterable[str], rows: Iterable[Iterable[Any]]) -> str:
    payload = json.dumps(
        {"title": title, "columns": list(columns), "rows": list(rows)},
        ensure_ascii=False,
        sort_keys=True,
        default=str,
        separators=(",", ":"),
    )
    return f"{_slug(title)}-{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:10]}"


def asset_hash(asset: StructuredAsset | dict[str, Any]) -> str:
    value = asset.model_dump(mode="json") if hasattr(asset, "model_dump") else asset
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _key(label: Any, index: int, used: set[str]) -> str:
    base = re.sub(r"\W+", "_", str(label or "").strip().lower(), flags=re.UNICODE).strip("_")
    base = base or f"column_{index + 1}"
    candidate, suffix = base, 2
    while candidate in used:
        candidate = f"{base}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _scalar(value: Any) -> Scalar:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return None
        if stripped.lower() in {"true", "false"}:
            return stripped.lower() == "true"
        if re.fullmatch(r"[-+]?(?:0|[1-9]\d*)(?:[.,]\d+)?", stripped):
            normalized = stripped.replace(",", ".")
            return float(normalized) if "." in normalized else int(normalized)
        return stripped
    if value is None or isinstance(value, (int, float, bool)):
        return value
    return str(value)


def _infer_type(values: list[Scalar]) -> str:
    present = [value for value in values if value is not None and value != ""]
    if not present:
        return "text"
    if all(isinstance(value, bool) for value in present):
        return "boolean"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in present):
        return "number"
    date_pattern = re.compile(r"^\d{4}-\d{1,2}-\d{1,2}(?:[T ].*)?$")
    if all(isinstance(value, str) and date_pattern.match(value.strip()) for value in present):
        return "date"
    return "text"


def dataset_from_matrix(
    title: str,
    matrix: list[list[Any]],
    *,
    source_refs: list[str] | None = None,
    asset_id: str | None = None,
    provenance: str = "extracted",
) -> DatasetAsset:
    rows = [list(row) for row in matrix if any(value not in (None, "") for value in row)]
    if not rows:
        raise StructuredAssetError(f"{title}: no tabular data")
    width = max(len(row) for row in rows)
    rows = [row + [None] * (width - len(row)) for row in rows]
    labels = [str(value).strip() if value not in (None, "") else f"Column {index + 1}" for index, value in enumerate(rows[0])]
    used: set[str] = set()
    keys = [_key(label, index, used) for index, label in enumerate(labels)]
    body = [[_scalar(value) for value in row] for row in rows[1:]]
    columns = [
        DatasetColumn(
            key=key,
            label=label,
            type=_infer_type([row[index] for row in body]),
            unit=(
                "%" if any(token in label.casefold() for token in ("percent", "percentage", "процент", "доля", "%"))
                else None
            ),
        )
        for index, (key, label) in enumerate(zip(keys, labels))
    ]
    return DatasetAsset(
        id=asset_id or _stable_id(title, labels, body),
        title=title,
        columns=columns,
        rows=[dict(zip(keys, row)) for row in body],
        source_refs=source_refs or [],
        provenance=provenance,
    )


def parse_csv_bytes(data: bytes, title: str = "CSV") -> list[DatasetAsset]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise StructuredAssetError("CSV must use UTF-8 or UTF-8 BOM") from exc
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    matrix = [row for row in csv.reader(io.StringIO(text), dialect) if any(cell.strip() for cell in row)]
    return [dataset_from_matrix(title, matrix)]


def parse_xlsx(path_or_stream: str | Path | io.BytesIO, title: str | None = None) -> list[DatasetAsset]:
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path_or_stream, read_only=True, data_only=True)
    except Exception as exc:
        raise StructuredAssetError("XLSX is damaged or unsupported") from exc
    assets: list[DatasetAsset] = []
    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
            matrix = [row for row in matrix if any(value not in (None, "") for value in row)]
            if matrix:
                assets.append(dataset_from_matrix(sheet.title if title is None else f"{title}: {sheet.title}", matrix))
    finally:
        workbook.close()
    if not assets:
        raise StructuredAssetError("XLSX has no visible non-empty sheets")
    return assets


def parse_json_assets(data: bytes | str) -> list[StructuredAsset]:
    try:
        raw = json.loads(data.decode("utf-8-sig") if isinstance(data, bytes) else data)
        if isinstance(raw, dict) and "structured_assets" in raw:
            raw = raw["structured_assets"]
        if isinstance(raw, dict):
            raw = [raw]
        return ASSET_LIST_ADAPTER.validate_python(raw)
    except (UnicodeDecodeError, json.JSONDecodeError, ValidationError, TypeError) as exc:
        raise StructuredAssetError(f"invalid structured assets JSON: {exc}") from exc


_MD_SEPARATOR = re.compile(r"^\s*:?-{3,}:?\s*$")


def parse_markdown_tables(text: str, title: str = "Markdown") -> list[DatasetAsset]:
    lines = text.replace("\r\n", "\n").splitlines()
    assets: list[DatasetAsset] = []
    index = 0
    while index + 1 < len(lines):
        header = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        separator = [cell.strip() for cell in lines[index + 1].strip().strip("|").split("|")]
        if len(header) >= 2 and len(separator) == len(header) and all(_MD_SEPARATOR.match(cell) for cell in separator):
            matrix: list[list[Any]] = [header]
            index += 2
            while index < len(lines) and "|" in lines[index]:
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                matrix.append((row + [""] * len(header))[:len(header)])
                index += 1
            assets.append(dataset_from_matrix(f"{title} table {len(assets) + 1}", matrix))
            continue
        index += 1
    return assets


def parse_docx_tables(path_or_stream: str | Path | io.BytesIO, title: str = "DOCX") -> list[DatasetAsset]:
    try:
        from docx import Document
        document = Document(path_or_stream)
    except Exception as exc:
        raise StructuredAssetError("DOCX is damaged or unsupported") from exc
    return [
        dataset_from_matrix(
            f"{title} table {index}",
            [[cell.text.strip() for cell in row.cells] for row in table.rows],
        )
        for index, table in enumerate(document.tables, 1)
        if table.rows
    ]


def _supports(
    capabilities: VisualCapabilities | None,
    kind: str,
    subtype: str | None,
) -> bool:
    if capabilities is None:
        return True
    return kind in capabilities.kinds and (
        not subtype or not capabilities.subtypes or subtype in capabilities.subtypes
    )


def choose_visual(
    asset: StructuredAsset,
    capabilities: VisualCapabilities | None = None,
    selected_columns: list[str] | None = None,
) -> tuple[str, str | None] | None:
    hint = asset.visual_hint
    if hint.mode == "none":
        return None
    if hint.mode == "force":
        subtype = hint.subtype
        if hint.kind in SUPPORTED_VISUAL_SUBTYPES and subtype not in SUPPORTED_VISUAL_SUBTYPES[hint.kind]:
            raise StructuredAssetError(f"unsupported forced {hint.kind} subtype {subtype!r}")
        forced = (hint.kind or "table", subtype)
        if not _supports(capabilities, *forced):
            raise StructuredAssetError(
                f"forced visual {forced[0]}/{forced[1] or '-'} is incompatible with slot capabilities"
            )
        return forced
    if isinstance(asset, DiagramAsset):
        preferred = ("diagram", asset.diagram_type)
        return preferred if _supports(capabilities, *preferred) else None
    if isinstance(asset, PictogramGridAsset):
        preferred = ("pictogram_grid", None)
        return preferred if _supports(capabilities, *preferred) else None
    selected = set(selected_columns or [])
    columns = [column for column in asset.columns if not selected or column.key in selected]
    numeric = [column for column in columns if column.type == "number"]
    dates = [column for column in columns if column.type == "date"]
    text = [column for column in columns if column.type == "text"]
    temporal_names = {
        "date", "datetime", "time", "day", "week", "month", "quarter", "year",
        "дата", "время", "день", "неделя", "месяц", "квартал", "год",
    }
    temporal = dates or [
        column for column in text
        if {part for part in re.split(r"[^\wа-яё]+", f"{column.key} {column.label}".casefold()) if part}
        & temporal_names
    ]
    percentage = [
        column for column in numeric
        if column.unit == "%" or any(
            token in f"{column.key} {column.label}".casefold()
            for token in ("percent", "percentage", "процент", "доля", "%")
        )
    ]
    if temporal and numeric:
        preferred = ("chart", "line")
    elif len(percentage) == 1 and text:
        preferred = ("chart", "donut")
    elif len(numeric) == 2 and not text:
        preferred = ("chart", "scatter")
    elif len(numeric) == 1 and len(text) == 1:
        labels = [str(row[text[0].key] or "") for row in asset.rows]
        if len(asset.rows) <= 6:
            preferred = ("chart", "donut")
        else:
            preferred = ("chart", "bar" if labels and sum(map(len, labels)) / len(labels) > 12 else "column")
    elif numeric and text and len(text) < len(columns):
        labels = [str(row[text[0].key] or "") for row in asset.rows]
        preferred = ("chart", "bar" if labels and sum(map(len, labels)) / len(labels) > 12 else "column")
    else:
        preferred = ("table", None)
    if _supports(capabilities, *preferred):
        return preferred
    fallback = ("table", None)
    return fallback if _supports(capabilities, *fallback) else None


def default_visual_columns(
    asset: StructuredAsset,
    visual_kind: str,
    subtype: str | None,
) -> list[str]:
    """Return a lossless default, except where a chart has a semantic series."""
    if not isinstance(asset, DatasetAsset):
        return []
    if visual_kind == "table":
        return [column.key for column in asset.columns]
    text = [column for column in asset.columns if column.type in {"text", "date"}]
    numeric = [column for column in asset.columns if column.type == "number"]
    if subtype in {"pie", "donut"}:
        percentage = [
            column for column in numeric
            if column.unit == "%" or any(
                token in f"{column.key} {column.label}".casefold()
                for token in ("percent", "percentage", "процент", "доля", "%")
            )
        ]
        if text and percentage:
            return [text[0].key, percentage[0].key]
    return [column.key for column in asset.columns]


def paginate_asset(
    asset: StructuredAsset,
    visual_kind: str,
    subtype: str | None,
    capabilities: VisualCapabilities | None,
    *,
    max_slides: int = 30,
    selected_columns: list[str] | None = None,
) -> list[VisualSlotContent]:
    max_items = (capabilities.max_items if capabilities else None) or 12
    selections: list[VisualSlotContent] = []
    if isinstance(asset, DatasetAsset):
        known_keys = [column.key for column in asset.columns]
        keys = list(selected_columns or known_keys)
        if not keys or len(keys) != len(set(keys)) or any(key not in known_keys for key in keys):
            raise StructuredAssetError("visual assignment selects invalid columns")
        if visual_kind == "table":
            max_rows = (capabilities.max_rows if capabilities else None) or 12
            max_columns = (capabilities.max_columns if capabilities else None) or len(keys)
            if max_columns < 1:
                raise StructuredAssetError("content_overflow")
            if len(keys) <= max_columns:
                blocks = [keys]
            elif max_columns == 1:
                blocks = [[key] for key in keys]
            else:
                blocks = [[keys[0], *keys[index:index + max_columns - 1]] for index in range(1, len(keys), max_columns - 1)]
            pages = max(1, (len(asset.rows) + max_rows - 1) // max_rows)
            selections = [
                VisualSlotContent(asset_id=asset.id, visual_kind="table", selected_columns=block, page=page)
                for block in blocks for page in range(1, pages + 1)
            ]
        else:
            pages = max(1, (len(asset.rows) + max_items - 1) // max_items)
            selections = [
                VisualSlotContent(asset_id=asset.id, visual_kind="chart", subtype=subtype, selected_columns=keys, page=page)
                for page in range(1, pages + 1)
            ]
    elif isinstance(asset, DiagramAsset):
        ordered = asset.order or [node.id for node in asset.nodes]
        selections = [
            VisualSlotContent(
                asset_id=asset.id, visual_kind="diagram", subtype=subtype or asset.diagram_type,
                selected_nodes=ordered[index:index + max_items], page=index // max_items + 1,
            )
            for index in range(0, len(ordered), max_items)
        ]
    else:
        selections = [
            VisualSlotContent(asset_id=asset.id, visual_kind="pictogram_grid", page=index // max_items + 1)
            for index in range(0, len(asset.items), max_items)
        ] or [VisualSlotContent(asset_id=asset.id, visual_kind="pictogram_grid")]
    if len(selections) > max_slides:
        raise StructuredAssetError("content_overflow")
    return selections


def materialize_visual(
    asset: StructuredAsset,
    selection: VisualSlotContent,
    capabilities: VisualCapabilities | None = None,
) -> dict[str, Any]:
    if asset.id != selection.asset_id:
        raise StructuredAssetError("visual assignment references a different asset")
    if isinstance(asset, DatasetAsset):
        keys = selection.selected_columns or [column.key for column in asset.columns]
        known = {column.key: column for column in asset.columns}
        if any(key not in known for key in keys):
            raise StructuredAssetError("visual assignment selects an unknown column")
        if selection.visual_kind == "table":
            page_size = (capabilities.max_rows if capabilities else None) or len(asset.rows) or 1
            start = (selection.page - 1) * page_size
            selected_rows = asset.rows[start:start + page_size]
            return {
                "title": asset.title,
                "columns": [known[key].label for key in keys],
                "rows": [[row[key] for key in keys] for row in selected_rows],
            }
        if selection.visual_kind != "chart":
            raise StructuredAssetError("dataset supports only table or chart visualization")
        number_keys = [key for key in keys if known[key].type == "number"]
        category_keys = [key for key in keys if known[key].type != "number"]
        subtype = selection.subtype or choose_visual(asset)[1]  # type: ignore[index]
        page_size = (capabilities.max_items if capabilities else None) or len(asset.rows) or 1
        start = (selection.page - 1) * page_size
        selected_rows = asset.rows[start:start + page_size]
        if subtype == "scatter":
            if len(number_keys) < 2:
                raise StructuredAssetError("scatter requires two numeric columns")
            categories, series_keys = [row[number_keys[0]] for row in selected_rows], [number_keys[1]]
        else:
            if not number_keys:
                raise StructuredAssetError("chart requires a numeric column")
            categories = [row[category_keys[0]] for row in selected_rows] if category_keys else list(range(start + 1, start + len(selected_rows) + 1))
            series_keys = number_keys
        all_values = [row[key] for row in asset.rows for key in series_keys if isinstance(row[key], (int, float)) and not isinstance(row[key], bool)]
        return {
            "title": asset.title,
            "chart_type": subtype,
            "categories": categories,
            "series": [{"name": known[key].label, "values": [row[key] for row in selected_rows]} for key in series_keys],
            "unit": known[series_keys[0]].unit if series_keys else None,
            "value_axis_min": min(all_values) if all_values else None,
            "value_axis_max": max(all_values) if all_values else None,
        }
    if isinstance(asset, DiagramAsset):
        selected = set(selection.selected_nodes or [node.id for node in asset.nodes])
        return {
            "title": asset.title,
            "diagram_type": selection.subtype or asset.diagram_type,
            "nodes": [node.model_dump(mode="json") for node in asset.nodes if node.id in selected],
            "edges": [edge.model_dump(mode="json") for edge in asset.edges if edge.source in selected and edge.target in selected],
            "order": [node_id for node_id in (asset.order or [node.id for node in asset.nodes]) if node_id in selected],
        }
    page_size = (capabilities.max_items if capabilities else None) or len(asset.items) or 1
    start = (selection.page - 1) * page_size
    return {"title": asset.title, "items": [item.model_dump(mode="json") for item in asset.items[start:start + page_size]]}
