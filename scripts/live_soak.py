from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
import json
import statistics
import tempfile
import time
from pathlib import Path

from pptx import Presentation

from content_planner.client import OpenAICompatibleClient
from content_planner.models import PlanningRequest, ProvidedImage
from content_planner.planner import ContentPlanner
from presentation_builder import PresentationBuilder


ROOT = Path(__file__).resolve().parents[1]
SCENARIO = ROOT / "description" / "tasks_examples" / "it_team_products"
TEMPLATE = ROOT / "template_model" / "output" / "run_5"


class CountingLLM:
    def __init__(self) -> None:
        self.delegate = OpenAICompatibleClient()
        self.model = self.delegate.model
        self.calls = 0

    def complete(self, prompt: str, **kwargs):
        self.calls += 1
        try:
            return self.delegate.complete(prompt, **kwargs)
        except Exception as exc:
            print(json.dumps({
                "llm_error_type": type(exc).__name__,
                "status": getattr(exc, "status_code", None),
                "body": getattr(exc, "body", None),
                "message": str(exc)[:500],
            }, ensure_ascii=False, default=str), flush=True)
            raise


def request() -> PlanningRequest:
    portraits = sorted((SCENARIO / "portraits").glob("*.png"))
    return PlanningRequest.model_validate({
        "brief": (SCENARIO / "brief.txt").read_text(encoding="utf-8"),
        "content_package": (SCENARIO / "content_package.md").read_text(encoding="utf-8"),
        "provided_images": [ProvidedImage(
            id=f"image:{path.stem}", original_name=path.name, asset_ref=path.name,
        ).model_dump(mode="json") for path in portraits],
        "slide_count": {"min": 10, "max": 10},
    })


def run_once(root: Path, number: int, *, strict: bool = False) -> dict[str, object]:
    started = time.monotonic()
    llm = CountingLLM()
    run_dir = root / f"run-{number:02d}"
    run_dir.mkdir(exist_ok=True)
    planner = ContentPlanner(
        TEMPLATE, llm=llm, enable_critique=strict, max_revisions=1,
        single_critique=False, allow_model_assets=False,
        semantic_pipeline=True,
        planning_budget=135, checkpoint_dir=run_dir / "checkpoints",
        job_id=f"soak-{number:02d}",
    )
    plan = planner.generate(request())
    plan_path = run_dir / "plan.json"
    plan_path.write_text(plan.model_dump_json(indent=2), encoding="utf-8")
    output = run_dir / "presentation.pptx"
    report = PresentationBuilder().build(
        TEMPLATE, plan, output, assets_dir=SCENARIO / "portraits", qa_dir=run_dir / "qa",
    )
    reopened = Presentation(output)
    elapsed = time.monotonic() - started
    fallback_used = plan.quality.result_kind == "fallback"
    result = {
        "run": number,
        "seconds": round(elapsed, 3),
        "llm_calls": llm.calls,
        "quality": plan.quality.status,
        "degraded": False,
        "slides": len(reopened.slides),
        "preview_warnings": sum(issue.code == "preview_failed" for issue in report.issues),
        "ok": (
            len(reopened.slides) == len(plan.slides)
            and llm.calls <= 30
            and elapsed <= 180
            and not fallback_used
            and not any(issue.level == "error" for issue in report.issues)
        ),
    }
    print(json.dumps(result, ensure_ascii=False), flush=True)
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--runs", type=int, default=20)
    parser.add_argument("--workers", type=int, default=1)
    parser.add_argument("--keep", type=Path)
    parser.add_argument("--strict", action="store_true")
    args = parser.parse_args()
    if args.runs < 1:
        raise ValueError("runs must be positive")
    if args.keep:
        args.keep.mkdir(parents=True, exist_ok=True)
        root_context = None
        root = args.keep
    else:
        root_context = tempfile.TemporaryDirectory(prefix="presd-soak-")
        root = Path(root_context.name)
    try:
        if args.workers < 1:
            raise ValueError("workers must be positive")
        with ThreadPoolExecutor(max_workers=min(args.workers, args.runs)) as executor:
            futures = {
                executor.submit(run_once, root, number, strict=args.strict): number
                for number in range(1, args.runs + 1)
            }
            results = [future.result() for future in as_completed(futures)]
        durations = sorted(float(item["seconds"]) for item in results)
        p95_index = max(0, min(len(durations) - 1, int(len(durations) * .95 + .999) - 1))
        summary = {
            "runs": len(results),
            "passed": sum(bool(item["ok"]) for item in results),
            "p95_seconds": durations[p95_index],
            "mean_seconds": round(statistics.mean(durations), 3),
            "max_llm_calls": max(int(item["llm_calls"]) for item in results),
        }
        print(json.dumps({"summary": summary}, ensure_ascii=False), flush=True)
        return 0 if summary["passed"] == args.runs and summary["p95_seconds"] <= 180 else 1
    finally:
        if root_context is not None:
            root_context.cleanup()


if __name__ == "__main__":
    raise SystemExit(main())
