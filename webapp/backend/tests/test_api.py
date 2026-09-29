from io import BytesIO
import json

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

from webapp.backend.app.database import SessionLocal
from webapp.backend.app.config import settings
from webapp.backend.app import api as api_module
from webapp.backend.app.extractors import extract_document, validate_content_file
from webapp.backend.app.models import Job, JobEvent, Template, User
from webapp.backend.app.storage import from_data, relative_to_data


ORIGIN = {"Origin": "http://testserver"}


def pptx_bytes() -> bytes:
    stream = BytesIO()
    Presentation().save(stream)
    return stream.getvalue()


def test_allowlisted_login_session_logout_and_origin(client):
    blocked = client.post("/api/v1/auth/login", json={"username": "alice", "password": "correct horse"}, headers={"Origin": "http://evil.invalid"})
    assert blocked.status_code == 403
    assert client.post("/api/v1/auth/register", json={"username": "mallory", "password": "correct horse"}, headers=ORIGIN).status_code == 404
    assert client.post("/api/v1/auth/login", json={"username": "mallory", "password": "correct horse"}, headers=ORIGIN).status_code == 401
    response = client.post("/api/v1/auth/login", json={"username": "alice", "password": "correct horse"}, headers=ORIGIN)
    assert response.status_code == 200
    assert response.cookies.get("presentation_session")
    assert client.get("/api/v1/auth/me").json()["username"] == "alice"
    assert client.post("/api/v1/auth/logout", headers=ORIGIN).status_code == 204
    assert client.get("/api/v1/auth/me").status_code == 401


def test_removing_client_from_allowlist_revokes_existing_session(client, monkeypatch):
    response = client.post(
        "/api/v1/auth/login",
        json={"username": "alice", "password": "correct horse"},
        headers=ORIGIN,
    )
    assert response.status_code == 200
    monkeypatch.setattr(settings, "allowed_clients_json", '{"bob":"another correct horse"}')
    assert client.get("/api/v1/auth/me").status_code == 401


def test_template_validation_deduplication_and_uuid_path(registered, clean_database):
    upload = {"file": ("../../brand.pptx", pptx_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    first = registered.post(
        "/api/v1/templates", files=upload,
        data={"use_vlm": "false", "catalog_workers": "2", "enrichment_workers": "3"},
        headers=ORIGIN,
    )
    assert first.status_code == 202
    second = registered.post("/api/v1/templates", files=upload, headers=ORIGIN)
    assert second.status_code == 202
    assert first.json()["id"] == second.json()["id"]
    assert first.json()["analysis_job"]["progress"] == 0
    assert first.json()["analysis_job"]["created_at"]
    assert first.json()["analysis_job"]["updated_at"]
    assert len(clean_database.calls) == 1
    assert clean_database.calls[0][1][1:] == (False, 2, False)
    with SessionLocal() as db:
        template = db.get(Template, first.json()["id"])
        job = db.get(Job, first.json()["analysis_job"]["id"])
        assert "brand.pptx" not in template.model_path
        assert from_data(template.model_path).parent.joinpath("source.pptx").is_file()
        request = json.loads(
            (from_data(job.artifact_path) / "request.json").read_text(encoding="utf-8")
        )
    assert request["enrichment_workers"] == 3


def test_template_upload_uses_server_enrichment_workers_default(registered, clean_database):
    response = registered.post(
        "/api/v1/templates",
        files={"file": ("legacy.pptx", pptx_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        headers=ORIGIN,
    )

    assert response.status_code == 202
    with SessionLocal() as db:
        job = db.get(Job, response.json()["analysis_job"]["id"])
        request = json.loads(
            (from_data(job.artifact_path) / "request.json").read_text(encoding="utf-8")
        )
    assert request["enrichment_workers"] == settings.template_enrichment_workers


@pytest.mark.parametrize("value", [0, 9])
def test_template_upload_validates_enrichment_worker_bounds(registered, value):
    response = registered.post(
        "/api/v1/templates",
        files={"file": ("brand.pptx", pptx_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        data={"enrichment_workers": str(value)}, headers=ORIGIN,
    )

    assert response.status_code == 422


def test_invalid_pptx_has_unified_error(registered):
    response = registered.post("/api/v1/templates", files={"file": ("bad.pptx", b"not a zip", "application/octet-stream")}, headers=ORIGIN)
    assert response.status_code == 422
    assert response.json()["code"] == "invalid_pptx"
    assert "message" in response.json()


@pytest.mark.parametrize("template_status", ["ready", "failed"])
def test_template_deletion_hides_record_and_removes_only_template_files(
    registered, template_status,
):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        template = Template(
            owner_id=user.id, name="brand.pptx", sha256=template_status[0] * 64,
            status=template_status,
        )
        db.add(template)
        db.flush()
        template_root = settings.data_root / "users" / user.id / "templates" / template.id
        template_root.mkdir(parents=True, exist_ok=True)
        (template_root / "source.pptx").write_bytes(b"source")
        template.model_path = relative_to_data(template_root / "model")
        job_root = settings.data_root / "users" / user.id / "jobs" / "completed"
        job_root.mkdir(parents=True, exist_ok=True)
        (job_root / "presentation.pptx").write_bytes(b"result")
        generation = Job(
            owner_id=user.id, template_id=template.id, type="generation",
            status="ready", stage="ready", brief_excerpt="Quarterly",
            artifact_path=relative_to_data(job_root),
        )
        db.add(generation)
        db.commit()
        template_id, generation_id = template.id, generation.id

    response = registered.delete(f"/api/v1/templates/{template_id}", headers=ORIGIN)

    assert response.status_code == 204
    assert not template_root.exists()
    assert job_root.is_dir()
    assert registered.get("/api/v1/templates").json()["items"] == []
    assert registered.get(f"/api/v1/templates/{template_id}").status_code == 404
    history = registered.get("/api/v1/jobs").json()["items"]
    assert history[0]["template_name"] == "brand.pptx"
    assert registered.get(f"/api/v1/jobs/{generation_id}/download").content == b"result"
    assert registered.delete(f"/api/v1/templates/{template_id}", headers=ORIGIN).status_code == 404
    with SessionLocal() as db:
        assert db.get(Template, template_id).status == "deleted"


def test_pdf_download_serves_existing_and_converts_legacy_job(registered, monkeypatch):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        root = settings.data_root / "users" / user.id / "jobs" / "pdf-download"
        root.mkdir(parents=True, exist_ok=True)
        (root / "presentation.pptx").write_bytes(b"pptx")
        job = Job(
            owner_id=user.id, type="generation", status="ready", stage="ready",
            artifact_path=relative_to_data(root),
        )
        db.add(job); db.commit(); job_id = job.id

    conversions: list[tuple[Path, Path]] = []

    def convert(source: Path, destination: Path) -> Path:
        conversions.append((source, destination))
        destination.write_bytes(b"%PDF-generated")
        return destination

    monkeypatch.setattr(api_module, "render_pdf", convert)

    response = registered.get(f"/api/v1/jobs/{job_id}/download/pdf")
    assert response.status_code == 200
    assert response.content == b"%PDF-generated"
    assert response.headers["content-type"] == "application/pdf"
    assert f'filename="presentation-{job_id}.pdf"' in response.headers["content-disposition"]
    assert conversions == [(
        (root / "presentation.pptx").resolve(),
        (root / "presentation.pdf").resolve(),
    )]

    assert registered.get(f"/api/v1/jobs/{job_id}/download/pdf").content == b"%PDF-generated"
    assert len(conversions) == 1


def test_pdf_conversion_failure_is_retryable_and_preserves_pptx(registered, monkeypatch):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        root = settings.data_root / "users" / user.id / "jobs" / "pdf-failure"
        root.mkdir(parents=True, exist_ok=True)
        source = root / "presentation.pptx"
        source.write_bytes(b"pptx")
        job = Job(
            owner_id=user.id, type="generation", status="ready", stage="ready",
            artifact_path=relative_to_data(root),
        )
        db.add(job); db.commit(); job_id = job.id

    def fail(*_args):
        raise RuntimeError("converter failed")

    monkeypatch.setattr(api_module, "render_pdf", fail)
    response = registered.get(f"/api/v1/jobs/{job_id}/download/pdf")

    assert response.status_code == 503
    assert response.json()["code"] == "pdf_unavailable"
    assert source.read_bytes() == b"pptx"
    assert not (root / "presentation.pdf").exists()


@pytest.mark.parametrize("job_type", ["template_analysis", "generation"])
@pytest.mark.parametrize("job_status", ["queued", "processing"])
def test_template_deletion_is_blocked_by_active_jobs(
    registered, job_type, job_status,
):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        template = Template(
            owner_id=user.id, name="busy.pptx",
            sha256=f"{job_type}-{job_status}".encode().hex().ljust(64, "0")[:64],
            status="processing" if job_type == "template_analysis" else "ready",
        )
        db.add(template)
        db.flush()
        root = settings.data_root / "users" / user.id / "templates" / template.id
        root.mkdir(parents=True, exist_ok=True)
        (root / "source.pptx").write_bytes(b"source")
        db.add(Job(
            owner_id=user.id, template_id=template.id, type=job_type,
            status=job_status,
        ))
        db.commit()
        template_id = template.id

    response = registered.delete(f"/api/v1/templates/{template_id}", headers=ORIGIN)

    assert response.status_code == 409
    assert response.json()["code"] == "template_in_use"
    assert root.is_dir()
    with SessionLocal() as db:
        assert db.get(Template, template_id).status != "deleted"


def test_foreign_template_cannot_be_deleted(registered, client):
    with SessionLocal() as db:
        alice = db.query(User).filter_by(username="alice").one()
        template = Template(
            owner_id=alice.id, name="private.pptx", sha256="f" * 64,
            status="ready",
        )
        db.add(template)
        db.commit()
        template_id = template.id
    client.post("/api/v1/auth/logout", headers=ORIGIN)
    client.post(
        "/api/v1/auth/login",
        json={"username": "bob", "password": "another correct horse"}, headers=ORIGIN,
    )

    assert client.delete(f"/api/v1/templates/{template_id}", headers=ORIGIN).status_code == 404


def test_reupload_reactivates_deleted_template_and_starts_fresh_analysis(
    registered, clean_database,
):
    data = pptx_bytes()
    digest = api_module.validate_pptx(data)
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        template = Template(
            owner_id=user.id, name="old-name.pptx", sha256=digest,
            status="deleted", slide_count=10, model_path="stale/model", error="old",
        )
        db.add(template)
        db.commit()
        template_id = template.id

    response = registered.post(
        "/api/v1/templates",
        files={"file": ("new-name.pptx", data, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        data={"use_vlm": "false"}, headers=ORIGIN,
    )

    assert response.status_code == 202
    assert response.json()["id"] == template_id
    assert response.json()["name"] == "new-name.pptx"
    assert response.json()["status"] == "processing"
    assert response.json()["analysis_job"]["status"] == "queued"
    assert len(clean_database.calls) == 1
    with SessionLocal() as db:
        template = db.get(Template, template_id)
        assert template.slide_count is None and template.error is None
        assert from_data(template.model_path).parent.joinpath("source.pptx").read_bytes() == data


def test_foreign_objects_are_hidden(client):
    client.post("/api/v1/auth/login", json={"username": "alice", "password": "correct horse"}, headers=ORIGIN)
    with SessionLocal() as db:
        alice = db.query(User).filter_by(username="alice").one()
        template = Template(owner_id=alice.id, name="private.pptx", sha256="a" * 64, status="ready", model_path="users/x/model")
        db.add(template); db.flush()
        job = Job(owner_id=alice.id, template_id=template.id, type="generation")
        db.add(job); db.commit()
        template_id, job_id = template.id, job.id
    client.post("/api/v1/auth/logout", headers=ORIGIN)
    client.post("/api/v1/auth/login", json={"username": "bob", "password": "another correct horse"}, headers=ORIGIN)
    assert client.get(f"/api/v1/templates/{template_id}").status_code == 404
    assert client.get(f"/api/v1/jobs/{job_id}").status_code == 404
    assert client.post(f"/api/v1/jobs/{job_id}/cancel", headers=ORIGIN).status_code == 404
    assert client.get(f"/api/v1/jobs/{job_id}/download").status_code == 404
    assert client.get(f"/api/v1/jobs/{job_id}/download/pdf").status_code == 404


def test_content_file_validation_and_extraction(tmp_path):
    image = BytesIO(); Image.new("RGB", (4, 4), "red").save(image, format="PNG")
    assert validate_content_file("photo.png", image.getvalue()) == ".png"
    document = Document(); document.add_paragraph("Первый абзац"); table = document.add_table(rows=1, cols=2); table.cell(0,0).text="A"; table.cell(0,1).text="B"
    path = tmp_path / "source.docx"; document.save(path)
    extracted = extract_document(path, "source.docx")
    assert "Первый абзац" in extracted and "A | B" in extracted


def test_job_form_validation(registered):
    response = registered.post("/api/v1/jobs", data={"template_id":"missing", "brief":"x", "content_text":"facts"}, headers=ORIGIN)
    assert response.status_code == 404


def test_cancel_active_generation_is_persisted_and_idempotent(registered, monkeypatch):
    stopped: list[str] = []
    monkeypatch.setattr(api_module, "_stop_rq_generation", stopped.append)
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        job = Job(
            owner_id=user.id, type="generation", status="processing",
            stage="planning", progress=42, brief_excerpt="Итоги",
        )
        db.add(job); db.commit(); job_id = job.id

    response = registered.post(f"/api/v1/jobs/{job_id}/cancel", headers=ORIGIN)

    assert response.status_code == 200
    assert response.json()["status"] == "canceled"
    assert response.json()["stage"] == "canceled"
    assert response.json()["progress"] == 42
    assert stopped == [job_id]
    with SessionLocal() as db:
        event = db.query(JobEvent).filter_by(job_id=job_id).one()
        assert event.stage == "canceled"
        assert event.message == "Задача отменена пользователем"

    repeated = registered.post(f"/api/v1/jobs/{job_id}/cancel", headers=ORIGIN)
    assert repeated.status_code == 200
    assert repeated.json()["status"] == "canceled"
    assert stopped == [job_id]


def test_cancel_rejects_terminal_and_non_generation_jobs(registered, monkeypatch):
    monkeypatch.setattr(api_module, "_stop_rq_generation", lambda _: None)
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        ready = Job(owner_id=user.id, type="generation", status="ready", stage="ready")
        analysis = Job(owner_id=user.id, type="template_analysis", status="processing")
        db.add_all([ready, analysis]); db.commit()
        ready_id, analysis_id = ready.id, analysis.id

    ready_response = registered.post(f"/api/v1/jobs/{ready_id}/cancel", headers=ORIGIN)
    analysis_response = registered.post(f"/api/v1/jobs/{analysis_id}/cancel", headers=ORIGIN)

    assert ready_response.status_code == 409
    assert ready_response.json()["code"] == "job_not_active"
    assert analysis_response.status_code == 409
    assert analysis_response.json()["code"] == "job_not_cancelable"


def test_rq_cancel_helper_cancels_queued_and_stops_started(monkeypatch):
    connection = object()
    monkeypatch.setattr(api_module.Redis, "from_url", lambda _: connection)
    stopped: list[tuple[object, str]] = []
    monkeypatch.setattr(api_module, "send_stop_job_command", lambda conn, job_id: stopped.append((conn, job_id)))

    class FakeRQJob:
        def __init__(self, status: str):
            self.status = status
            self.cancel_calls = 0

        def get_status(self, refresh: bool = False):
            assert refresh is True
            return self.status

        def cancel(self):
            self.cancel_calls += 1

    queued = FakeRQJob("queued")
    monkeypatch.setattr(api_module.RQJob, "fetch", lambda job_id, connection: queued)
    api_module._stop_rq_generation("queued-id")
    assert queued.cancel_calls == 1

    started = FakeRQJob("started")
    monkeypatch.setattr(api_module.RQJob, "fetch", lambda job_id, connection: started)
    api_module._stop_rq_generation("started-id")
    assert stopped == [(connection, "generation-started-id")]


def test_reanalysis_rejects_active_generation_and_passes_settings(registered, clean_database, tmp_path):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        model_dir = from_data(f"users/{user.id}/test-model")
        model_dir.mkdir(parents=True, exist_ok=True)
        template = Template(owner_id=user.id, name="brand.pptx", sha256="b" * 64, status="ready", model_path=relative_to_data(model_dir))
        db.add(template); db.flush()
        generation = Job(owner_id=user.id, template_id=template.id, type="generation", status="processing")
        db.add(generation); db.commit()
        template_id, generation_id = template.id, generation.id
    blocked = registered.post(f"/api/v1/templates/{template_id}/reanalyze", json={"use_vlm":False,"catalog_workers":2}, headers=ORIGIN)
    assert blocked.status_code == 409
    with SessionLocal() as db:
        generation = db.get(Job, generation_id)
        generation.status = "ready"
        db.commit()
    response = registered.post(
        f"/api/v1/templates/{template_id}/reanalyze",
        json={"use_vlm": False, "catalog_workers": 2, "enrichment_workers": 5},
        headers=ORIGIN,
    )
    assert response.status_code == 202
    assert response.json()["analysis_job"]["status"] == "queued"
    assert response.json()["analysis_job"]["created_at"]
    assert response.json()["analysis_job"]["updated_at"]
    assert clean_database.calls[-1][1][1:] == (False, 2, True)
    with SessionLocal() as db:
        job = db.get(Job, response.json()["analysis_job"]["id"])
        request = json.loads(
            (from_data(job.artifact_path) / "request.json").read_text(encoding="utf-8")
        )
    assert request["enrichment_workers"] == 5


@pytest.mark.parametrize("value", [0, 9])
def test_template_reanalysis_validates_enrichment_worker_bounds(registered, value):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        template = Template(
            owner_id=user.id, name=f"bounds-{value}.pptx", sha256=str(value) * 64,
            status="ready", model_path=f"users/{user.id}/bounds-{value}",
        )
        db.add(template)
        db.commit()
        template_id = template.id

    response = registered.post(
        f"/api/v1/templates/{template_id}/reanalyze",
        json={"enrichment_workers": value}, headers=ORIGIN,
    )

    assert response.status_code == 422


def test_generation_fast_mode_is_persisted(registered, clean_database):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        template = Template(owner_id=user.id, name="brand.pptx", sha256="c" * 64, status="ready", model_path=f"users/{user.id}/model")
        db.add(template); db.commit(); template_id = template.id
    response = registered.post("/api/v1/jobs", data={"template_id":template_id,"brief":"Итоги","content_text":"Факты","fast_mode":"true"}, headers=ORIGIN)
    assert response.status_code == 202
    with SessionLocal() as db:
        job = db.get(Job, response.json()["job_id"])
        request_path = from_data(job.artifact_path) / "request.json"
    request = json.loads(request_path.read_text(encoding="utf-8"))
    assert request["fast_mode"] is True
    assert request["model_path"] == f"users/{user.id}/model"


def test_reanalysis_allows_active_generation_with_pinned_model(registered, clean_database):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        root = settings.data_root / "users" / user.id
        old_model = root / "old-model"
        old_model.mkdir(parents=True, exist_ok=True)
        template = Template(
            owner_id=user.id, name="brand.pptx", sha256="7" * 64,
            status="ready", model_path=relative_to_data(old_model),
        )
        db.add(template); db.flush()
        generation_root = root / "jobs" / "active-pinned"
        generation_root.mkdir(parents=True, exist_ok=True)
        (generation_root / "request.json").write_text(json.dumps({
            "model_path": relative_to_data(old_model),
        }), encoding="utf-8")
        generation = Job(
            owner_id=user.id, template_id=template.id, type="generation",
            status="processing", artifact_path=relative_to_data(generation_root),
        )
        db.add(generation); db.commit(); template_id = template.id

    response = registered.post(
        f"/api/v1/templates/{template_id}/reanalyze",
        json={"use_vlm": False, "catalog_workers": 2}, headers=ORIGIN,
    )

    assert response.status_code == 202
    assert response.json()["status"] == "ready"
    assert response.json()["analysis_job"]["status"] == "queued"


def test_generation_mode_defaults_to_strict_and_explicit_value_wins(registered, clean_database):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        template = Template(
            owner_id=user.id, name="brand.pptx", sha256="9" * 64,
            status="ready", model_path=f"users/{user.id}/model-9",
        )
        db.add(template); db.commit(); template_id = template.id
    first = registered.post(
        "/api/v1/jobs",
        data={"template_id": template_id, "brief": "Итоги", "content_text": "Факты"},
        headers=ORIGIN,
    )
    assert first.status_code == 202
    second = registered.post(
        "/api/v1/jobs",
        data={
            "template_id": template_id, "brief": "Итоги", "content_text": "Факты",
            "fast_mode": "true", "generation_mode": "strict",
        },
        headers=ORIGIN,
    )
    assert second.status_code == 202
    with SessionLocal() as db:
        first_job = db.get(Job, first.json()["job_id"])
        second_job = db.get(Job, second.json()["job_id"])
        first_request = json.loads((from_data(first_job.artifact_path) / "request.json").read_text())
        second_request = json.loads((from_data(second_job.artifact_path) / "request.json").read_text())
    assert first_request["generation_mode"] == "strict"
    assert second_request["generation_mode"] == "strict"
    assert clean_database.calls[-1][2]["job_id"].startswith("generation-")
    assert ":" not in clean_database.calls[-1][2]["job_id"]


def test_generation_accepts_csv_and_structured_assets_json(registered, clean_database):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        template = Template(owner_id=user.id, name="visual.pptx", sha256="d" * 64, status="ready", model_path=f"users/{user.id}/visual-model")
        db.add(template); db.commit(); template_id = template.id
    structured = '{"id":"manual","kind":"dataset","title":"Manual","columns":[{"key":"x","label":"X","type":"number"}],"rows":[{"x":1}]}'
    response = registered.post(
        "/api/v1/jobs",
        data={"template_id": template_id, "brief": "Итоги", "structured_assets": structured},
        files={"files": ("revenue.csv", b"Month;Revenue\nJan;120", "text/csv")},
        headers=ORIGIN,
    )
    assert response.status_code == 202
    with SessionLocal() as db:
        job = db.get(Job, response.json()["job_id"])
        request = (from_data(job.artifact_path) / "request.json").read_text(encoding="utf-8")
    assert '"id": "manual"' in request


def test_generation_rejects_invalid_structured_assets(registered, clean_database):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        template = Template(owner_id=user.id, name="visual.pptx", sha256="e" * 64, status="ready", model_path=f"users/{user.id}/visual-model")
        db.add(template); db.commit(); template_id = template.id
    response = registered.post(
        "/api/v1/jobs",
        data={"template_id": template_id, "brief": "Итоги", "structured_assets": '{"kind":"dataset"}'},
        headers=ORIGIN,
    )
    assert response.status_code == 422
    assert response.json()["code"] == "invalid_structured_assets"


def test_failed_job_exposes_structured_validation_issues(registered, clean_database):
    with SessionLocal() as db:
        user = db.query(User).filter_by(username="alice").one()
        template = Template(
            owner_id=user.id, name="visual.pptx", sha256="f" * 64,
            status="ready", model_path=f"users/{user.id}/visual-model",
        )
        db.add(template); db.flush()
        job = Job(
            owner_id=user.id, template_id=template.id, type="generation",
            status="failed", stage="failed", error="Данные не помещаются в диапазон.",
            warnings=[{
                "code": "content_overflow", "path": "structured_assets[0]",
                "message": "visual pagination requires 11 slides; maximum is 10",
            }],
        )
        db.add(job); db.commit(); job_id = job.id
    response = registered.get(f"/api/v1/jobs/{job_id}")
    assert response.status_code == 200
    payload = response.json()
    assert payload["warnings"] == []
    assert payload["validation_issues"][0]["code"] == "content_overflow"


def test_storage_rejects_traversal():
    try:
        from_data("../../etc/passwd")
    except Exception as exc:
        assert getattr(exc, "code", None) == "not_found"
    else:
        raise AssertionError("path traversal was accepted")
