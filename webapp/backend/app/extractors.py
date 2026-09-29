from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader

from content_planner.models import StructuredAsset
from content_planner.structured_assets import (
    StructuredAssetError,
    parse_csv_bytes,
    parse_docx_tables,
    parse_json_assets,
    parse_markdown_tables,
    parse_xlsx,
)


TEXT_EXTENSIONS = {".txt", ".md"}
STRUCTURED_EXTENSIONS = {".csv", ".xlsx", ".json"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
ALLOWED_EXTENSIONS = TEXT_EXTENSIONS | STRUCTURED_EXTENSIONS | IMAGE_EXTENSIONS | {".pdf", ".docx"}


class ContentError(ValueError):
    pass


def validate_content_file(name: str, data: bytes) -> str:
    extension = Path(name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ContentError("Допустимы TXT, MD, PDF, DOCX, CSV, XLSX, JSON, PNG и JPEG")
    try:
        if extension in IMAGE_EXTENSIONS:
            with Image.open(BytesIO(data)) as image:
                image.verify()
        elif extension == ".pdf":
            PdfReader(BytesIO(data))
        elif extension == ".docx":
            Document(BytesIO(data))
        elif extension in TEXT_EXTENSIONS:
            data.decode("utf-8")
        elif extension == ".csv":
            parse_csv_bytes(data, Path(name).stem)
        elif extension == ".xlsx":
            parse_xlsx(BytesIO(data), Path(name).stem)
        elif extension == ".json":
            parse_json_assets(data)
    except (UnicodeDecodeError, OSError, ValueError, KeyError, StructuredAssetError) as exc:
        raise ContentError(f"Файл {name} повреждён или имеет неверный формат") from exc
    return extension


def extract_document(path: Path, original_name: str) -> str:
    extension = Path(original_name).suffix.lower()
    if extension in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8")
    if extension == ".pdf":
        reader = PdfReader(path)
        pages = []
        for number, page in enumerate(reader.pages, 1):
            value = (page.extract_text() or "").strip()
            if value:
                pages.append(f"[PDF, страница {number}]\n{value}")
        return "\n\n".join(pages)
    if extension == ".docx":
        document = Document(path)
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        # Keep a readable representation for the narrative model as well as
        # the lossless dataset returned by extract_structured_assets().
        for table in document.tables:
            blocks.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(blocks)
    return ""


def extract_structured_assets(path: Path, original_name: str) -> list[StructuredAsset]:
    extension = Path(original_name).suffix.lower()
    title = Path(original_name).stem
    try:
        if extension == ".csv":
            return parse_csv_bytes(path.read_bytes(), title)
        if extension == ".xlsx":
            return parse_xlsx(path, title)
        if extension == ".json":
            return parse_json_assets(path.read_bytes())
        if extension == ".md":
            return parse_markdown_tables(path.read_text(encoding="utf-8"), title)
        if extension == ".docx":
            return parse_docx_tables(path, title)
    except StructuredAssetError as exc:
        raise ContentError(str(exc)) from exc
    return []
