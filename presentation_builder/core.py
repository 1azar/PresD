from __future__ import annotations

import hashlib
import io
import json
import os
import shutil
import tempfile
from copy import deepcopy
from collections.abc import Callable
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR, MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml

from content_planner.models import (
    CATALOG_SCHEMA_VERSION,
    PLAN_SCHEMA_VERSION,
    BindingRenderer,
    PresentationPlan,
    RenderOperation,
    SlotAction,
    TemplateCatalog,
)
from content_planner.compiler import CompilationError, compile_slide
from content_planner.structured_assets import asset_hash
from template_model.analyzer import render_previews

from .models import BuildFailure, BuildIssue, BuildReport, ImageGenerator
from .icons import resolve_icon


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _string_values(value: Any):
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for child in value.values():
            yield from _string_values(child)
    elif isinstance(value, list):
        for child in value:
            yield from _string_values(child)


def _content_count(content: Any) -> int:
    data = content.model_dump(mode="python")
    if content.kind in {"bullet_list", "cards", "profiles", "metrics", "timeline", "process"}:
        return len(data["items"])
    if content.kind == "table":
        return len(data["rows"])
    if content.kind == "chart":
        return len(data["categories"])
    return 1


def _replace_relationship_ids(element: Any, mapping: dict[str, str]) -> None:
    relationship_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    for node in element.iter():
        for name, value in list(node.attrib.items()):
            if name.startswith(f"{{{relationship_namespace}}}") and value in mapping:
                node.set(name, mapping[value])


def _clone_slide(presentation: Presentation, source: Any) -> Any:
    destination = presentation.slides.add_slide(source.slide_layout)
    destination_element = destination._element
    destination_element.clear()
    destination_element.attrib.update(source._element.attrib)
    for child in source._element:
        destination_element.append(deepcopy(child))
    destination.shapes._element = destination_element.spTree
    destination.shapes._spTree = destination_element.spTree
    destination.shapes._grpSp = destination_element.spTree
    destination.shapes._cached_max_shape_id = None

    mapping: dict[str, str] = {}
    destination_layout_rel = next(
        rel for rel in destination.part.rels.values() if rel.reltype == RT.SLIDE_LAYOUT
    )
    for rel in source.part.rels.values():
        if rel.reltype == RT.SLIDE_LAYOUT:
            mapping[rel.rId] = destination_layout_rel.rId
            continue
        if rel.reltype == RT.NOTES_SLIDE:
            continue
        if rel.is_external:
            new_rid = destination.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = destination.part.relate_to(rel.target_part, rel.reltype)
        mapping[rel.rId] = new_rid
    _replace_relationship_ids(destination_element, mapping)
    return destination


def _remove_original_slides(presentation: Presentation, count: int) -> None:
    slide_ids = presentation.slides._sldIdLst
    for slide_id in list(slide_ids)[:count]:
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def _walk_shapes(shapes: Any):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)


def _shape_by_id(slide: Any, shape_id: int) -> Any | None:
    return next((shape for shape in _walk_shapes(slide.shapes) if shape.shape_id == shape_id), None)


def _set_paragraph_text(paragraph: Any, value: str) -> None:
    runs = list(paragraph.runs)
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run().text = value


def _replace_text(shape: Any, paragraphs: list[str]) -> None:
    if not getattr(shape, "has_text_frame", hasattr(shape, "text_frame")):
        raise ValueError("target shape has no text frame")
    frame = shape.text_frame
    existing = list(frame.paragraphs)
    while len(existing) < len(paragraphs):
        prototype = existing[-1]._p if existing else frame.paragraphs[0]._p
        clone = deepcopy(prototype)
        frame._txBody.append(clone)
        existing = list(frame.paragraphs)
    for index, value in enumerate(paragraphs):
        _set_paragraph_text(existing[index], value)
    for paragraph in existing[len(paragraphs):]:
        frame._txBody.remove(paragraph._p)
    if not paragraphs:
        _set_paragraph_text(frame.paragraphs[0], "")


def _resolve_asset(reference: str, assets_dir: Path) -> Path:
    candidate = (assets_dir / reference).resolve()
    try:
        candidate.relative_to(assets_dir)
    except ValueError as exc:
        raise ValueError("asset_ref escapes assets_dir") from exc
    if not candidate.is_file():
        raise ValueError(f"asset not found: {reference}")
    return candidate


def _crop_picture(shape: Any, image_path: Path, fit: str) -> None:
    if fit == "stretch":
        return
    with Image.open(image_path) as image:
        image_ratio = image.width / image.height
    box_ratio = shape.width / shape.height
    shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
    if fit == "cover":
        if image_ratio > box_ratio:
            crop = (1 - box_ratio / image_ratio) / 2
            shape.crop_left = shape.crop_right = crop
        else:
            crop = (1 - image_ratio / box_ratio) / 2
            shape.crop_top = shape.crop_bottom = crop


def _replace_picture(slide: Any, shape: Any, path: Path, operation: RenderOperation) -> None:
    if shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
        _, r_id = slide.part.get_or_add_image_part(str(path))
        shape._pic.blipFill.blip.embed = r_id
        _crop_picture(shape, path, operation.image_fit)
        if operation.alt_text:
            shape._pic.nvPicPr.cNvPr.set("descr", operation.alt_text)
        return

    parent = shape._element.getparent()
    if parent is not slide.shapes._spTree:
        raise ValueError("image_box inside a group is not supported")
    index = parent.index(shape._element)
    picture = slide.shapes.add_picture(str(path), shape.left, shape.top, shape.width, shape.height)
    picture_element = picture._element
    picture_element.getparent().remove(picture_element)
    parent.remove(shape._element)
    parent.insert(index, picture_element)
    _crop_picture(picture, path, operation.image_fit)
    if operation.alt_text:
        picture._pic.nvPicPr.cNvPr.set("descr", operation.alt_text)


def _replace_table(shape: Any, data: dict[str, Any]) -> None:
    if not getattr(shape, "has_table", False):
        raise ValueError("target shape is not a table")
    columns = data.get("columns", [])
    rows = data.get("rows", [])
    values = [columns, *rows]
    if len(values) > len(shape.table.rows) or any(len(row) > len(shape.table.columns) for row in values):
        raise ValueError("table data exceeds the template grid")
    for row_index, row in enumerate(shape.table.rows):
        for column_index, cell in enumerate(row.cells):
            value = values[row_index][column_index] if row_index < len(values) and column_index < len(values[row_index]) else ""
            _replace_text(cell, ["" if value is None else str(value)])


def _replace_chart(shape: Any, data: dict[str, Any]) -> None:
    if not getattr(shape, "has_chart", False):
        raise ValueError("target shape is not a chart")
    if data.get("chart_type") == "scatter":
        chart_data = XyChartData()
        categories = data.get("categories", [])
        for item in data.get("series", []):
            series = chart_data.add_series(item["name"])
            for x, y in zip(categories, item["values"]):
                series.add_data_point(float(x), float(y))
    else:
        chart_data = CategoryChartData()
        chart_data.categories = ["" if value is None else str(value) for value in data.get("categories", [])]
        for item in data.get("series", []):
            chart_data.add_series(item["name"], item["values"])
    shape.chart.replace_data(chart_data)
    _set_chart_scale(shape.chart, data)


def _set_chart_scale(chart: Any, data: dict[str, Any]) -> None:
    try:
        if data.get("value_axis_min") is not None:
            chart.value_axis.minimum_scale = float(data["value_axis_min"])
        if data.get("value_axis_max") is not None:
            chart.value_axis.maximum_scale = float(data["value_axis_max"])
    except (AttributeError, ValueError, TypeError):
        pass


_CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE,
    "donut": XL_CHART_TYPE.DOUGHNUT,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}


def _move_to_anchor_z(slide: Any, shape: Any, anchor: Any) -> None:
    parent = anchor._element.getparent()
    if parent is not slide.shapes._spTree:
        raise ValueError("visual anchor inside a group is not supported")
    index = parent.index(anchor._element)
    element = shape._element
    element.getparent().remove(element)
    parent.insert(index, element)


def _remove_anchor(slide: Any, anchor: Any) -> int:
    parent = anchor._element.getparent()
    if parent is not slide.shapes._spTree:
        raise ValueError("visual anchor inside a group is not supported")
    index = parent.index(anchor._element)
    parent.remove(anchor._element)
    return index


def _create_table(slide: Any, anchor: Any, data: dict[str, Any]) -> Any:
    columns = data.get("columns") or []
    rows = data.get("rows") or []
    if not columns:
        raise ValueError("table visualization has no columns")
    shape = slide.shapes.add_table(
        max(1, len(rows) + 1), len(columns), anchor.left, anchor.top, anchor.width, anchor.height
    )
    _move_to_anchor_z(slide, shape, anchor)
    values = [columns, *rows]
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            shape.table.cell(row_index, column_index).text = "" if value is None else str(value)
    _remove_anchor(slide, anchor)
    return shape


def _chart_data(data: dict[str, Any]) -> Any:
    if data.get("chart_type") == "scatter":
        result = XyChartData()
        for item in data.get("series", []):
            series = result.add_series(item["name"])
            for x, y in zip(data.get("categories", []), item.get("values", [])):
                series.add_data_point(float(x), float(y))
        return result
    result = CategoryChartData()
    result.categories = ["" if value is None else str(value) for value in data.get("categories", [])]
    for item in data.get("series", []):
        result.add_series(item["name"], item.get("values", []))
    return result


def _create_chart(slide: Any, anchor: Any, data: dict[str, Any]) -> Any:
    subtype = data.get("chart_type")
    if subtype not in _CHART_TYPES:
        raise ValueError(f"unsupported chart type: {subtype}")
    shape = slide.shapes.add_chart(
        _CHART_TYPES[subtype], anchor.left, anchor.top, anchor.width, anchor.height, _chart_data(data)
    )
    _set_chart_scale(shape.chart, data)
    shape.chart.has_title = False
    shape.chart.has_legend = (
        len(data.get("series", [])) > 1 or subtype in {"pie", "donut"}
    )
    _move_to_anchor_z(slide, shape, anchor)
    _remove_anchor(slide, anchor)
    return shape


def _clear_chart_canvas(slide: Any, preserve_shape_ids: list[int]) -> None:
    preserve = set(preserve_shape_ids)

    def prune(shapes: Any) -> bool:
        """Remove canvas content bottom-up while retaining ancestors of kept shapes."""
        kept_any = False
        for shape in reversed(list(shapes)):
            keep = shape.shape_id in preserve
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                keep = prune(shape.shapes) or keep
            if keep:
                kept_any = True
                continue
            parent = shape._element.getparent()
            if parent is not None:
                parent.remove(shape._element)
        return kept_any

    prune(slide.shapes)


def _create_chart_on_canvas(
    slide: Any, operation: RenderOperation, slide_width: int, slide_height: int,
) -> Any:
    bounds = operation.canvas_bounds
    data = operation.data or {}
    subtype = data.get("chart_type")
    if bounds is None or subtype not in _CHART_TYPES:
        raise ValueError(f"unsupported chart canvas subtype: {subtype}")
    shape = slide.shapes.add_chart(
        _CHART_TYPES[subtype],
        int(bounds.x * slide_width), int(bounds.y * slide_height),
        int(bounds.width * slide_width), int(bounds.height * slide_height),
        _chart_data(data),
    )
    _set_chart_scale(shape.chart, data)
    shape.chart.has_title = False
    shape.chart.has_legend = (
        len(data.get("series", [])) > 1 or subtype in {"pie", "donut"}
    )
    return shape


def _set_node_style(shape: Any, text: str, accent: RGBColor = RGBColor(48, 99, 219)) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = accent
    shape.line.color.rgb = accent
    shape.text = text
    frame = shape.text_frame
    frame.word_wrap = True
    for paragraph in frame.paragraphs:
        paragraph.alignment = 1
        for run in paragraph.runs:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(255, 255, 255)


def _diagram_positions(data: dict[str, Any], left: int, top: int, width: int, height: int) -> dict[str, tuple[int, int, int, int]]:
    nodes = data.get("nodes", [])
    count = max(1, len(nodes))
    subtype = data.get("diagram_type", "process")
    positions: dict[str, tuple[int, int, int, int]] = {}
    if subtype == "hierarchy":
        by_id = {node["id"]: node for node in nodes}
        levels: list[list[dict[str, Any]]] = []
        remaining = list(nodes)
        known: set[str] = set()
        while remaining:
            level = [node for node in remaining if not node.get("parent_id") or node.get("parent_id") in known]
            if not level:
                level = remaining
            levels.append(level)
            known.update(node["id"] for node in level)
            remaining = [node for node in remaining if node["id"] not in known]
        node_h = max(1, int(height / max(1, len(levels)) * 0.55))
        for level_index, level in enumerate(levels):
            node_w = max(1, int(width / max(1, len(level)) * 0.72))
            for index, node in enumerate(level):
                x = left + int((index + 0.5) * width / len(level) - node_w / 2)
                y = top + int((level_index + 0.5) * height / len(levels) - node_h / 2)
                positions[node["id"]] = (x, y, node_w, node_h)
    elif subtype == "cycle":
        import math
        node_w, node_h = int(width * 0.24), int(height * 0.18)
        for index, node in enumerate(nodes):
            angle = -math.pi / 2 + 2 * math.pi * index / count
            x = left + int(width / 2 + math.cos(angle) * width * 0.34 - node_w / 2)
            y = top + int(height / 2 + math.sin(angle) * height * 0.34 - node_h / 2)
            positions[node["id"]] = (x, y, node_w, node_h)
    else:
        vertical = subtype == "timeline" and height > width
        if vertical:
            node_w, node_h = int(width * 0.72), max(1, int(height / count * 0.62))
            for index, node in enumerate(nodes):
                positions[node["id"]] = (left + int(width * 0.14), top + int((index + 0.18) * height / count), node_w, node_h)
        else:
            node_w, node_h = max(1, int(width / count * 0.72)), int(height * 0.42)
            for index, node in enumerate(nodes):
                positions[node["id"]] = (left + int((index + 0.14) * width / count), top + int(height * 0.29), node_w, node_h)
    return positions


def _create_diagram(slide: Any, anchor: Any, data: dict[str, Any]) -> list[Any]:
    positions = _diagram_positions(data, anchor.left, anchor.top, anchor.width, anchor.height)
    nodes = {node["id"]: node for node in data.get("nodes", [])}
    if not nodes:
        raise ValueError("diagram visualization has no nodes")
    created: list[Any] = []
    for node_id in data.get("order") or nodes:
        node = nodes[node_id]
        x, y, width, height = positions[node_id]
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, width, height)
        value = node["label"]
        if node.get("description"):
            value += f"\n{node['description']}"
        _set_node_style(shape, value)
        created.append(shape)
        if node.get("icon_query"):
            icon_path, _ = resolve_icon(str(node["icon_query"]))
            icon_size = min(width, height) // 4
            created.append(_add_svg_picture(
                slide, icon_path, x + width - icon_size - max(1, width // 20),
                y + max(1, height // 12), icon_size, icon_size,
            ))
    edges = data.get("edges", [])
    if not edges:
        ordered = list(data.get("order") or nodes)
        edges = [{"source": first, "target": second} for first, second in zip(ordered, ordered[1:])]
    for edge in edges:
        if edge["source"] not in positions or edge["target"] not in positions:
            continue
        sx, sy, sw, sh = positions[edge["source"]]
        tx, ty, tw, th = positions[edge["target"]]
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, sx + sw // 2, sy + sh // 2, tx + tw // 2, ty + th // 2
        )
        connector.line.color.rgb = RGBColor(112, 124, 144)
        created.insert(0, connector)
    for shape in created:
        _move_to_anchor_z(slide, shape, anchor)
    _remove_anchor(slide, anchor)
    return created


def _add_svg_picture(slide: Any, svg_path: Path, left: int, top: int, width: int, height: int) -> Any:
    fallback = Image.new("RGBA", (128, 128), (255, 255, 255, 0))
    draw = ImageDraw.Draw(fallback)
    draw.ellipse((14, 14, 114, 114), outline=(48, 99, 219, 255), width=9)
    draw.line((42, 68, 61, 87, 92, 45), fill=(48, 99, 219, 255), width=9, joint="curve")
    stream = io.BytesIO()
    fallback.save(stream, format="PNG")
    stream.seek(0)
    picture = slide.shapes.add_picture(stream, left, top, width, height)
    picture._pic.nvPicPr.cNvPr.set("descr", svg_path.stem)

    package = slide.part.package
    partname = package.next_partname(PackURI("/ppt/media/presd-icon%d.svg"))
    svg = svg_path.read_text(encoding="utf-8").replace("currentColor", "#3063DB").encode("utf-8")
    svg_part = Part(partname, "image/svg+xml", package, svg)
    svg_rid = slide.part.relate_to(svg_part, RT.IMAGE)
    blip = picture._pic.blipFill.blip
    blip.append(parse_xml(
        '<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
        f'<asvg:svgBlip r:embed="{svg_rid}"/>'
        '</a:ext></a:extLst>'
    ))
    return picture


def _create_pictogram_grid(slide: Any, anchor: Any, data: dict[str, Any]) -> list[Any]:
    items = data.get("items", [])
    if not items:
        raise ValueError("pictogram grid has no items")
    import math
    columns = min(4, max(1, math.ceil(math.sqrt(len(items)))))
    rows = math.ceil(len(items) / columns)
    created: list[Any] = []
    for index, item in enumerate(items):
        column, row = index % columns, index // columns
        cell_w, cell_h = anchor.width / columns, anchor.height / rows
        x, y = int(anchor.left + column * cell_w), int(anchor.top + row * cell_h)
        icon_path, _ = resolve_icon(str(item.get("icon_query") or "unknown"))
        icon = _add_svg_picture(
            slide, icon_path, x + int(cell_w * 0.36), y + int(cell_h * 0.08),
            int(cell_w * 0.28), int(cell_h * 0.34),
        )
        label = slide.shapes.add_textbox(x + int(cell_w * 0.06), y + int(cell_h * 0.48), int(cell_w * 0.88), int(cell_h * 0.42))
        value = item.get("label", "")
        if item.get("value") is not None:
            value = f"{item['value']}\n{value}"
        label.text = value
        for paragraph in label.text_frame.paragraphs:
            paragraph.alignment = 1
        created.extend([icon, label])
    for shape in created:
        _move_to_anchor_z(slide, shape, anchor)
    _remove_anchor(slide, anchor)
    return created


def _create_visual(slide: Any, anchor: Any, operation: RenderOperation) -> Any:
    data = operation.data or {}
    if operation.visual_kind == "table":
        return _create_table(slide, anchor, data)
    if operation.visual_kind == "chart":
        return _create_chart(slide, anchor, data)
    if operation.visual_kind == "diagram":
        return _create_diagram(slide, anchor, data)
    if operation.visual_kind == "pictogram_grid":
        return _create_pictogram_grid(slide, anchor, data)
    raise ValueError(f"unsupported visual kind: {operation.visual_kind}")


class PresentationBuilder:
    def __init__(
        self,
        *,
        image_generator: ImageGenerator | None = None,
        libreoffice_path: str = "libreoffice",
        pdftoppm_path: str = "pdftoppm",
    ) -> None:
        self.image_generator = image_generator
        self.libreoffice_path = libreoffice_path
        self.pdftoppm_path = pdftoppm_path

    def build(
        self,
        template_dir: str | Path,
        plan: PresentationPlan | str | Path,
        output_path: str | Path,
        *,
        assets_dir: str | Path | None = None,
        qa_dir: str | Path | None = None,
        pdf_output_path: str | Path | None = None,
        progress: Callable[[str], None] | None = None,
    ) -> BuildReport:
        template_root = Path(template_dir).resolve()
        output = Path(output_path).resolve()
        assets = Path(assets_dir).resolve() if assets_dir else template_root
        qa = Path(qa_dir).resolve() if qa_dir else output.with_suffix(".qa")
        pdf_output = Path(pdf_output_path).resolve() if pdf_output_path else None
        if output.exists() or qa.exists() or (pdf_output is not None and pdf_output.exists()):
            raise BuildFailure("output_exists", f"output, PDF, or QA directory already exists: {output}, {pdf_output}, {qa}")
        if isinstance(plan, PresentationPlan):
            parsed = plan
        else:
            parsed = PresentationPlan.model_validate_json(Path(plan).read_text(encoding="utf-8"))
        manifest_path = template_root / "presentation.json"
        catalog_path = template_root / "planner_catalog.json"
        source_path = template_root / "source.pptx"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        try:
            catalog = TemplateCatalog.model_validate_json(catalog_path.read_text(encoding="utf-8"))
        except Exception as exc:
            raise BuildFailure("unsupported_catalog", f"invalid catalog: {exc}") from exc
        if parsed.schema_version != PLAN_SCHEMA_VERSION:
            raise BuildFailure("unsupported_plan", f"expected plan schema {PLAN_SCHEMA_VERSION}")
        actual_hash = _sha256(source_path)
        expected_hash = parsed.template.source_sha256
        if (
            actual_hash != expected_hash
            or manifest.get("source_sha256") != expected_hash
            or catalog.source_sha256 != expected_hash
        ):
            raise BuildFailure("template_mismatch", "plan and template source_sha256 do not match")
        if parsed.template.catalog_version != CATALOG_SCHEMA_VERSION:
            raise BuildFailure("catalog_mismatch", "plan references an incompatible catalog version")

        structured_assets = {asset.id: asset for asset in parsed.structured_assets}

        families = {family.family_id: family for family in catalog.families}
        for plan_slide in parsed.slides:
            family = families.get(plan_slide.template_family_id)
            if family is None or family.slide_class != plan_slide.template_class:
                raise BuildFailure("catalog_mismatch", f"slide {plan_slide.number} references an unknown family")
            variant = next(
                (item for item in family.variants if item.slide_number == plan_slide.template_slide_number),
                None,
            )
            if variant is None:
                raise BuildFailure("catalog_mismatch", f"slide {plan_slide.number} references an unknown variant")
            slots = {slot.slot_id: slot for slot in variant.slots}
            assignments = {assignment.slot_id: assignment for assignment in plan_slide.assignments}
            if set(assignments) != set(slots):
                raise BuildFailure(
                    "catalog_mismatch",
                    f"slide {plan_slide.number} assignments do not cover catalog slots exactly",
                )
            for slot_id, slot in slots.items():
                assignment = assignments[slot_id]
                if assignment.kind != slot.kind or assignment.target_shape_ids != slot.target_shape_ids:
                    raise BuildFailure("catalog_mismatch", f"slide {plan_slide.number} slot {slot_id!r} does not match catalog")
                if slot.required and assignment.action == SlotAction.CLEAR:
                    raise BuildFailure("required_slot_cleared", f"slide {plan_slide.number} slot {slot_id!r} is required")
                if assignment.content is None:
                    continue
                if slot.capacity is not None and _content_count(assignment.content) > slot.capacity:
                    raise BuildFailure("capacity_exceeded", f"slide {plan_slide.number} slot {slot_id!r} exceeds capacity {slot.capacity}")
                if slot.max_chars is not None and any(
                    len(value) > slot.max_chars
                    for value in _string_values(assignment.content.model_dump(mode="python"))
                ):
                    raise BuildFailure("text_too_long", f"slide {plan_slide.number} slot {slot_id!r} exceeds max_chars {slot.max_chars}")
            try:
                expected_operations = compile_slide(plan_slide, variant, structured_assets)
            except CompilationError as exc:
                raise BuildFailure("render_compilation_failed", str(exc)) from exc
            if expected_operations != plan_slide.render_operations:
                raise BuildFailure(
                    "render_operations_mismatch",
                    f"slide {plan_slide.number} render_operations do not match catalog bindings",
                )
            for operation in plan_slide.render_operations:
                if operation.asset_id:
                    asset = structured_assets.get(operation.asset_id)
                    if asset is None or operation.asset_hash != asset_hash(asset):
                        raise BuildFailure(
                            "structured_asset_mismatch",
                            f"slide {plan_slide.number} visual asset hash does not match the plan",
                        )

        presentation = Presentation(source_path)
        originals = list(presentation.slides)
        issues: list[BuildIssue] = []
        output.parent.mkdir(parents=True, exist_ok=True)
        qa.parent.mkdir(parents=True, exist_ok=True)
        generated_dir = Path(tempfile.mkdtemp(prefix=".generated-", dir=output.parent))
        staged_output: Path | None = None
        staged_qa: Path | None = None
        try:
            for plan_slide in parsed.slides:
                if plan_slide.assignments and not plan_slide.render_operations:
                    raise BuildFailure(
                        "missing_render_operations",
                        f"slide {plan_slide.number} has assignments but no compiled render operations",
                    )
                index = plan_slide.template_slide_number - 1
                if index < 0 or index >= len(originals):
                    raise BuildFailure("unknown_slide", f"template slide {plan_slide.template_slide_number} does not exist")
                slide = _clone_slide(presentation, originals[index])
                canvas_operations = [
                    operation for operation in plan_slide.render_operations
                    if operation.renderer == BindingRenderer.CANVAS
                ]
                if len(canvas_operations) > 1:
                    raise BuildFailure(
                        "render_operation_failed",
                        f"slide {plan_slide.number} contains multiple chart canvases",
                    )
                if canvas_operations:
                    _clear_chart_canvas(slide, canvas_operations[0].preserve_shape_ids)
                missing_shape_ids = sorted({
                    operation.shape_id
                    for operation in plan_slide.render_operations
                    if operation.renderer != BindingRenderer.CANVAS
                    and operation.shape_id is not None
                    and _shape_by_id(slide, operation.shape_id) is None
                })
                if missing_shape_ids:
                    raise BuildFailure(
                        "unknown_shape",
                        f"slide {plan_slide.number} is missing render targets after canvas cleanup: "
                        f"{missing_shape_ids}",
                    )
                for operation in [
                    *(
                        item for item in plan_slide.render_operations
                        if item.renderer != BindingRenderer.CANVAS
                    ),
                    *canvas_operations,
                ]:
                    if operation.visual_kind in {"pictogram_grid", "diagram"}:
                        icon_items = (
                            (operation.data or {}).get("items", [])
                            if operation.visual_kind == "pictogram_grid"
                            else (operation.data or {}).get("nodes", [])
                        )
                        for item in icon_items:
                            if operation.visual_kind == "diagram" and not item.get("icon_query"):
                                continue
                            _, matched = resolve_icon(str(item.get("icon_query") or ""))
                            if not matched:
                                issues.append(BuildIssue(
                                    level="warning",
                                    code="icon_fallback",
                                    message=f"No local icon match for {item.get('icon_query')!r}; neutral icon used",
                                    slide_number=plan_slide.number,
                                    shape_id=operation.shape_id,
                                ))
                    if operation.renderer == BindingRenderer.CANVAS:
                        try:
                            _create_chart_on_canvas(
                                slide, operation, presentation.slide_width, presentation.slide_height,
                            )
                        except (OSError, ValueError) as exc:
                            raise BuildFailure(
                                "render_operation_failed", str(exc),
                                [BuildIssue(
                                    level="error", code="render_operation_failed",
                                    message=str(exc), slide_number=plan_slide.number,
                                )],
                            ) from exc
                        continue
                    shape = _shape_by_id(slide, operation.shape_id)  # type: ignore[arg-type]
                    if shape is None:
                        raise BuildFailure(
                            "unknown_shape",
                            f"shape {operation.shape_id} does not exist on template slide {plan_slide.template_slide_number}",
                        )
                    try:
                        self._apply_operation(slide, shape, operation, assets, generated_dir, plan_slide.number)
                    except (OSError, ValueError) as exc:
                        raise BuildFailure(
                            "render_operation_failed",
                            str(exc),
                            [BuildIssue(level="error", code="render_operation_failed", message=str(exc), slide_number=plan_slide.number, shape_id=operation.shape_id)],
                        ) from exc
            _remove_original_slides(presentation, len(originals))
            if len(presentation.slides) != len(parsed.slides):
                raise BuildFailure("slide_count_mismatch", "assembled slide count is incorrect")

            handle, temporary_name = tempfile.mkstemp(prefix=f".{output.name}.", suffix=".pptx", dir=output.parent)
            os.close(handle)
            staged_output = Path(temporary_name)
            presentation.save(staged_output)
            reopened = Presentation(staged_output)
            if len(reopened.slides) != len(parsed.slides):
                raise BuildFailure("invalid_output", "saved presentation could not be verified")
            os.replace(staged_output, output)
            staged_output = None

            staged_qa = Path(tempfile.mkdtemp(prefix=f".{qa.name}.", dir=qa.parent))
            preview_dir = staged_qa / "previews"
            previews: list[Path] = []
            try:
                if progress is not None:
                    progress("rendering_preview")
                previews = render_previews(
                    output, preview_dir, self.libreoffice_path, self.pdftoppm_path,
                    pdf_output=pdf_output,
                )
                if len(previews) != len(parsed.slides):
                    issues.append(BuildIssue(level="warning", code="preview_count", message=f"renderer produced {len(previews)} previews for {len(parsed.slides)} slides"))
            except Exception as exc:
                issues.append(BuildIssue(level="warning", code="preview_failed", message=f"{type(exc).__name__}: {exc}"))
            if parsed.quality.status == "needs_review":
                issues.append(BuildIssue(
                    level="warning", code="plan_needs_review",
                    message="План презентации требует ручной проверки",
                ))
            report = BuildReport(
                output=str(output),
                slide_count=len(parsed.slides),
                qa_dir=str(qa),
                previews=[str(Path("previews") / path.name) for path in previews],
                issues=issues,
            )
            (staged_qa / "report.json").write_text(report.model_dump_json(indent=2) + "\n", encoding="utf-8")
            os.replace(staged_qa, qa)
            staged_qa = None
            return report
        finally:
            shutil.rmtree(generated_dir, ignore_errors=True)
            if staged_output and staged_output.exists():
                staged_output.unlink()
            if staged_qa and staged_qa.exists():
                shutil.rmtree(staged_qa, ignore_errors=True)

    def _apply_operation(
        self,
        slide: Any,
        shape: Any,
        operation: RenderOperation,
        assets: Path,
        generated_dir: Path,
        slide_number: int,
    ) -> None:
        if operation.action == SlotAction.KEEP:
            return
        if operation.action == SlotAction.CLEAR:
            if operation.renderer == BindingRenderer.TEXT:
                _replace_text(shape, [""])
            elif operation.renderer in {BindingRenderer.IMAGE, BindingRenderer.IMAGE_BOX}:
                shape._element.getparent().remove(shape._element)
            elif operation.renderer == BindingRenderer.TABLE:
                _replace_table(shape, {"columns": [], "rows": []})
            elif operation.renderer == BindingRenderer.VISUAL:
                shape._element.getparent().remove(shape._element)
            return
        if operation.renderer == BindingRenderer.TEXT:
            _replace_text(shape, operation.text or [""])
            return
        if operation.renderer == BindingRenderer.TABLE:
            _replace_table(shape, operation.data or {})
            return
        if operation.renderer == BindingRenderer.CHART:
            _replace_chart(shape, operation.data or {})
            return
        if operation.renderer == BindingRenderer.VISUAL:
            _create_visual(slide, shape, operation)
            return
        if operation.action == SlotAction.GENERATE:
            if self.image_generator is None:
                raise ValueError("image_generator_unavailable")
            generated = self.image_generator.generate(
                operation.prompt or "",
                alt_text=operation.alt_text,
                slide_number=slide_number,
                shape_id=operation.shape_id,
            )
            if isinstance(generated, bytes):
                path = generated_dir / f"slide_{slide_number:03d}_{operation.shape_id}.png"
                path.write_bytes(generated)
            else:
                path = Path(generated).resolve()
                if not path.is_file():
                    raise ValueError("image generator returned a missing file")
        else:
            path = _resolve_asset(operation.asset_ref or "", assets)
        _replace_picture(slide, shape, path, operation)
