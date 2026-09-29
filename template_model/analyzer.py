from __future__ import annotations

import hashlib
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import (
    Box,
    ChartContent,
    PresentationManifest,
    SlideElement,
    SlideManifestEntry,
    SlideMetadata,
    SlideVLMAnalysis,
    TableCell,
    TableContent,
    VLMStatus,
)
from .structures import detect_structures
from .vlm import VLMAnalyzer


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _box(shape: Any, width: int, height: int) -> Box:
    def bounded(value: float) -> float:
        return min(1.0, max(0.0, round(value, 6)))

    return Box(
        x=bounded(shape.left / width),
        y=bounded(shape.top / height),
        width=bounded(shape.width / width),
        height=bounded(shape.height / height),
    )


def _shape_text(shape: Any) -> str | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    text = shape.text
    return text if text.strip() else None


def _placeholder_type(shape: Any) -> str | None:
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        return str(shape.placeholder_format.type).split(" (")[0].lower()
    except (AttributeError, ValueError):
        return "unknown"


def _table_content(shape: Any) -> TableContent:
    cells = [
        TableCell(row=row_index, column=column_index, text=cell.text)
        for row_index, row in enumerate(shape.table.rows)
        for column_index, cell in enumerate(row.cells)
    ]
    return TableContent(
        rows=len(shape.table.rows),
        columns=len(shape.table.columns),
        cells=cells,
    )


def _chart_content(shape: Any) -> ChartContent:
    chart = shape.chart
    title: str | None = None
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text.strip() or None
    except (AttributeError, ValueError):
        pass

    series_names: list[str] = []
    try:
        for series in chart.series:
            name = str(series.name).strip()
            if name:
                series_names.append(name)
    except (AttributeError, ValueError, TypeError):
        pass

    category_labels: list[str] = []
    try:
        plots = list(chart.plots)
        if plots:
            for category in plots[0].categories:
                label = str(category.label).strip()
                if label:
                    category_labels.append(label)
    except (AttributeError, ValueError, TypeError):
        pass

    return ChartContent(
        title=title,
        chart_type=str(getattr(chart, "chart_type", "")) or None,
        category_labels=list(dict.fromkeys(category_labels)),
        series_names=list(dict.fromkeys(series_names)),
    )


def _alt_text(shape: Any) -> str | None:
    try:
        nodes = shape._element.xpath(".//p:cNvPr")
        if nodes:
            value = (nodes[0].get("descr") or nodes[0].get("title") or "").strip()
            return value or None
    except (AttributeError, IndexError, TypeError):
        pass
    return None


def _shape_type(shape: Any, text: str | None) -> str | None:
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_chart", False):
        return "chart"
    shape_type = shape.shape_type
    if shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
        return "image"
    if shape_type == MSO_SHAPE_TYPE.MEDIA:
        return "media"
    if text:
        return "text"
    if shape_type in {MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT}:
        return "other"
    marker = f"{getattr(shape, 'name', '')} {_alt_text(shape) or ''}".lower()
    return "visual_anchor" if "presd:visual" in marker else None


def extract_elements(slide: Any, width: int, height: int) -> list[SlideElement]:
    """Extract addressable content while omitting content-free decorative shapes."""

    def visit(shapes: Iterable[Any]) -> list[SlideElement]:
        result: list[SlideElement] = []
        for z_order, shape in enumerate(shapes):
            text = _shape_text(shape)
            element_type = _shape_type(shape, text)
            children = visit(shape.shapes) if element_type == "group" else []
            if element_type == "group" and not children:
                continue
            if element_type is None:
                continue
            result.append(SlideElement(
                shape_id=shape.shape_id,
                name=shape.name,
                type=element_type,
                box=_box(shape, width, height),
                z_order=z_order,
                placeholder_type=_placeholder_type(shape),
                alt_text=_alt_text(shape),
                text=text,
                table=_table_content(shape) if element_type == "table" else None,
                chart=_chart_content(shape) if element_type == "chart" else None,
                children=children,
            ))
        return result

    return visit(slide.shapes)


def _slide_style_tokens(slide: Any) -> tuple[list[str], list[str]]:
    fonts: set[str] = set()
    colors: set[str] = set()

    def visit(shapes: Iterable[Any]) -> None:
        for shape in shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts.add(run.font.name)
                        try:
                            rgb = run.font.color.rgb
                            if rgb is not None:
                                colors.add(str(rgb))
                        except (AttributeError, TypeError):
                            pass
            for owner in (getattr(shape, "fill", None), getattr(shape, "line", None)):
                try:
                    rgb = owner.fore_color.rgb if hasattr(owner, "fore_color") else owner.color.rgb
                    if rgb is not None:
                        colors.add(str(rgb))
                except (AttributeError, TypeError):
                    pass
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes)

    visit(slide.shapes)
    return sorted(fonts), sorted(colors)


def render_pdf(source: Path, destination: Path, libreoffice: str = "libreoffice") -> Path:
    source = Path(source).resolve()
    destination = Path(destination).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".pdf-", dir=destination.parent) as temporary:
        temp = Path(temporary)
        profile = temp / "libreoffice-profile"
        profile.mkdir()
        subprocess.run(
            [
                libreoffice,
                "--headless",
                "--nologo",
                "--nodefault",
                "--nofirststartwizard",
                f"-env:UserInstallation={profile.as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temp),
                str(source),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=30,
        )
        pdfs = [path for path in temp.glob("*.pdf") if path.is_file()]
        if len(pdfs) != 1 or pdfs[0].stat().st_size == 0:
            raise RuntimeError("LibreOffice did not produce a PDF")
        with pdfs[0].open("rb") as stream:
            if stream.read(5) != b"%PDF-":
                raise RuntimeError("LibreOffice produced an invalid PDF")
        os.replace(pdfs[0], destination)
    return destination


def render_previews(
    source: Path,
    destination: Path,
    libreoffice: str,
    pdftoppm: str,
    pdf_output: Path | None = None,
) -> list[Path]:
    destination.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".render-", dir=destination.parent) as temporary:
        temp = Path(temporary)
        pdf = render_pdf(source, pdf_output or temp / "presentation.pdf", libreoffice)
        prefix = temp / "slide"
        subprocess.run(
            [pdftoppm, "-png", "-r", "120", str(pdf), str(prefix)],
            check=True,
            capture_output=True,
            text=True,
            timeout=30,
        )
        generated = sorted(temp.glob("slide-*.png"))
        result: list[Path] = []
        for number, path in enumerate(generated, 1):
            target = destination / f"slide_{number:03d}.png"
            shutil.copy2(path, target)
            result.append(target)
        return result


def analyze_presentation(
    template_path: str | Path,
    staging_dir: str | Path,
    *,
    use_vlm: bool = False,
    vlm_client: Any | None = None,
    libreoffice: str = "libreoffice",
    pdftoppm: str = "pdftoppm",
) -> PresentationManifest:
    source = Path(template_path).resolve()
    staging = Path(staging_dir)
    if not source.is_file() or source.suffix.lower() != ".pptx":
        raise ValueError(f"Not a PPTX file: {source}")

    presentation = Presentation(source)
    source_hash = sha256_file(source)
    destination = staging / "source.pptx"
    shutil.copy2(source, destination)

    warnings: list[str] = []
    rendered_dir = staging / ".rendered_previews"
    previews: list[Path] = []
    try:
        previews = render_previews(destination, rendered_dir, libreoffice, pdftoppm)
        if len(previews) != len(presentation.slides):
            warnings.append(
                f"Renderer produced {len(previews)} previews for {len(presentation.slides)} slides"
            )
    except Exception as exc:
        warnings.append(f"Preview rendering failed: {type(exc).__name__}: {exc}")

    vlm: VLMAnalyzer | None = None
    vlm_error: str | None = None
    if use_vlm:
        if vlm_client is None and not VLMAnalyzer.credentials_available():
            vlm_error = "VLM endpoint and model are not configured"
        else:
            try:
                vlm = VLMAnalyzer(client=vlm_client)
            except Exception as exc:
                vlm_error = f"VLM initialization failed: {type(exc).__name__}: {exc}"
        if vlm_error:
            warnings.append(vlm_error)

    manifest_entries: list[SlideManifestEntry] = []
    vlm_failures = 0
    for index, slide in enumerate(presentation.slides, 1):
        slide_directory = staging / "slides" / f"slide_{index:03d}"
        slide_directory.mkdir(parents=True, exist_ok=True)
        elements = extract_elements(slide, presentation.slide_width, presentation.slide_height)

        preview_path: str | None = None
        if index <= len(previews):
            shutil.copy2(previews[index - 1], slide_directory / "preview.png")
            preview_path = "preview.png"

        slide_warnings: list[str] = []
        if not use_vlm:
            vlm_result = SlideVLMAnalysis(status=VLMStatus.DISABLED)
        elif preview_path is None:
            vlm_failures += 1
            vlm_result = SlideVLMAnalysis(
                status=VLMStatus.ERROR,
                error="Preview is unavailable; VLM analysis was skipped",
            )
        elif vlm is None:
            vlm_failures += 1
            vlm_result = SlideVLMAnalysis(status=VLMStatus.ERROR, error=vlm_error or "VLM is unavailable")
        else:
            vlm_structure = {
                "slide_number": index,
                "elements": [element.model_dump(mode="json") for element in elements],
            }
            try:
                result = vlm.analyze(slide_directory / "preview.png", vlm_structure)
                vlm_result = SlideVLMAnalysis(
                    status=VLMStatus.OK,
                    classification=result.classification,
                    tags=result.tags,
                    description=result.description,
                    confidence=result.confidence,
                )
            except Exception as exc:
                vlm_failures += 1
                error = f"{type(exc).__name__}: {exc}"
                slide_warnings.append(f"VLM analysis failed: {error}")
                vlm_result = SlideVLMAnalysis(status=VLMStatus.ERROR, error=error)

        layout = slide.slide_layout
        structures = detect_structures(
            elements,
            vlm_result.classification.value if vlm_result.classification else None,
        )
        theme_fonts, palette = _slide_style_tokens(slide)
        metadata = SlideMetadata(
            slide_number=index,
            slide_id=slide.slide_id,
            layout_name=getattr(layout, "name", None),
            preview_path=preview_path,
            elements=elements,
            structures=structures,
            theme_fonts=theme_fonts,
            palette=palette,
            vlm=vlm_result,
            warnings=slide_warnings,
        )
        metadata_path = slide_directory / "metadata.json"
        metadata_path.write_text(metadata.model_dump_json(indent=2), encoding="utf-8")

        relative_directory = f"slides/slide_{index:03d}"
        manifest_entries.append(SlideManifestEntry(
            slide_number=index,
            slide_id=slide.slide_id,
            directory=relative_directory,
            metadata_path=f"{relative_directory}/metadata.json",
            preview_path=f"{relative_directory}/preview.png" if preview_path else None,
            vlm_status=vlm_result.status,
        ))

    if use_vlm and vlm_failures:
        warnings.append(f"VLM analysis failed or was skipped for {vlm_failures} slide(s)")
    shutil.rmtree(rendered_dir, ignore_errors=True)

    return PresentationManifest(
        source_file="source.pptx",
        source_sha256=source_hash,
        slide_width=presentation.slide_width,
        slide_height=presentation.slide_height,
        slide_count=len(manifest_entries),
        vlm_enabled=use_vlm,
        slides=manifest_entries,
        warnings=warnings,
    )
