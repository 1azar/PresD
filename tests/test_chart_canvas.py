from __future__ import annotations

import hashlib
import json
import re
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from content_planner.catalog import TemplateCatalogBuilder, select_chart_canvas
from content_planner.compiler import compile_slide
from content_planner.models import (
    CATALOG_SCHEMA_VERSION,
    CatalogFamily,
    CatalogSlot,
    CatalogSlideDescriptor,
    CatalogVariant,
    ContentAnalysis,
    DeckInfo,
    PlanCandidate,
    PlanningRequest,
    PresentationPlan,
    QualityReport,
    SlidePlan,
    SourceChunk,
    SourcedText,
    TemplateCatalog,
    TemplateReference,
    VisualHint,
)
from content_planner.planner import ContentPlanner, PlannerFailure
from content_planner.structured_assets import parse_csv_bytes
from content_planner.validator import validate_candidate
from presentation_builder import PresentationBuilder


def canvas_slot(
    *, preferred: list[str], x: float = .03, y: float = .22,
    width: float = .94, height: float = .68, preserve: list[int] | None = None,
) -> CatalogSlot:
    return CatalogSlot.model_validate({
        "slot_id": "chart_visual", "kind": "visual", "role": "native chart canvas",
        "target_shape_ids": [], "bindings": [],
        "visual_capabilities": {
            "kinds": ["chart"],
            "subtypes": ["bar", "column", "line", "area", "pie", "donut", "scatter"],
            "render_mode": "container", "width": width, "height": height,
            "aspect_ratio": width / height, "max_items": 12,
        },
        "chart_canvas": {
            "bounds": {"x": x, "y": y, "width": width, "height": height},
            "preserve_shape_ids": preserve or [],
            "supported_subtypes": ["bar", "column", "line", "area", "pie", "donut", "scatter"],
            "preferred_subtypes": preferred,
        },
    })


def chart_catalog(specs: list[tuple[int, list[str], float]]) -> TemplateCatalog:
    families = []
    for number, preferred, width in specs:
        families.append(CatalogFamily(
            family_id=f"chart_{number}", slide_class="chart", description="Chart",
            variants=[CatalogVariant(
                slide_number=number, description="Chart",
                slots=[canvas_slot(preferred=preferred, width=width)], static_shape_ids=[],
            )],
        ))
    return TemplateCatalog(model="test", source_sha256="hash", families=families)


class ChartCanvasCatalogTests(unittest.TestCase):
    def test_selection_prefers_exact_then_family_then_area_and_slide_number(self):
        catalog = chart_catalog([
            (7, ["bar"], .80),
            (4, ["donut"], .90),
            (2, ["column"], .70),
        ])
        self.assertEqual(select_chart_canvas(catalog, "bar")[1].slide_number, 7)
        self.assertEqual(select_chart_canvas(catalog, "line")[1].slide_number, 7)
        self.assertEqual(select_chart_canvas(catalog, "pie")[1].slide_number, 4)

        tied = chart_catalog([(9, [], .90), (3, [], .90), (1, [], .70)])
        self.assertEqual(select_chart_canvas(tied, "scatter")[1].slide_number, 3)

    def test_builds_canvas_bounds_and_preserves_title_brand_and_footer(self):
        slide = {
            "slide_number": 1, "classification": "chart", "description": "Line chart",
            "tags": ["line"], "structures": [],
            "top_level_elements": [
                {"shape_id": 1, "type": "text", "placeholder_type": "title",
                 "name": "Title", "box": {"x": .04, "y": .05, "width": .8, "height": .12}},
                {"shape_id": 2, "type": "image", "name": "Brand logo",
                 "box": {"x": .9, "y": .03, "width": .05, "height": .05}},
                {"shape_id": 3, "type": "text", "name": "Footer",
                 "box": {"x": .03, "y": .93, "width": .2, "height": .03}},
                {"shape_id": 4, "type": "chart", "name": "Demo",
                 "box": {"x": .1, "y": .3, "width": .8, "height": .5}},
            ],
        }
        title = CatalogSlot.model_validate({
            "slot_id": "title", "kind": "text", "role": "title",
            "target_shape_ids": [1],
            "bindings": [{"shape_id": 1, "renderer": "text", "value_paths": ["/text"]}],
        })
        slots, static = TemplateCatalogBuilder._with_chart_canvas(
            slide, [title], [2, 3, 4], dedicated=True,
        )
        canvas = slots[-1].chart_canvas
        self.assertEqual(canvas.bounds.model_dump(), {"x": .03, "y": .22, "width": .94, "height": .68})
        self.assertEqual(canvas.preserve_shape_ids, [1, 2, 3])
        self.assertEqual(canvas.preferred_subtypes, ["line"])
        self.assertEqual(static, [2, 3])

    def test_slide_38_geometry_adds_fallback_title_without_ownership_overlap(self):
        slide = {
            "slide_number": 38, "classification": "content_visual",
            "description": "Chart layout", "tags": [], "structures": [],
            "elements": [
                {"shape_id": 2104, "type": "text", "placeholder_type": "title",
                 "text": "Заголовок в две или в одну строчку",
                 "box": {"x": .019463, "y": .055011, "width": .742890,
                         "height": .140800}},
                {"shape_id": 2105, "type": "image", "name": "Google Shape;2105;p79",
                 "box": {"x": .058333, "y": .295397, "width": .883329,
                         "height": .035181}},
                {"shape_id": 2106, "type": "image", "name": "Google Shape;2106;p79",
                 "box": {"x": .084375, "y": .260216, "width": .830725,
                         "height": .035181}},
                {"shape_id": 2107, "type": "image", "name": "Google Shape;2107;p79",
                 "box": {"x": .031250, "y": .330983, "width": .936974,
                         "height": .668952}},
            ],
        }
        slide["top_level_elements"] = list(slide["elements"])
        descriptor = CatalogSlideDescriptor.model_validate({
            "slide_number": 38, "classification": "content_visual",
            "description": "Chart layout", "family_hint": "charts",
            "family_description": "Charts", "slots": [{
                "slot_id": "hero", "kind": "image", "role": "visual",
                "required": False,
                "target": {"target_type": "shape", "shape_id": 2107},
            }],
        })
        builder = TemplateCatalogBuilder(Path("."), llm=object())
        base_slots, base_static = builder._expand_descriptor(slide, descriptor)

        fallback_title = next(slot for slot in base_slots if slot.slot_id == "fallback_title")
        self.assertEqual(fallback_title.target_shape_ids, [2104])
        self.assertFalse(fallback_title.required)
        self.assertEqual(fallback_title.bindings[0].value_paths, ["/text"])

        for dedicated in (True, False):
            slots, static = builder._with_chart_canvas(
                slide, base_slots, base_static, dedicated=dedicated,
            )
            variant = CatalogVariant(
                slide_number=38, description="Chart layout",
                slots=slots, static_shape_ids=static,
            )
            editable = {
                shape_id for slot in variant.slots for shape_id in slot.target_shape_ids
            }
            self.assertTrue(any(slot.chart_canvas is not None for slot in variant.slots))
            self.assertEqual(editable & set(variant.static_shape_ids), set())
            self.assertNotIn(2107, variant.static_shape_ids)

    def test_canvas_early_exits_preserve_original_slots_and_static_ids(self):
        visual = CatalogSlot.model_validate({
            "slot_id": "visual", "kind": "image", "role": "visual",
            "target_shape_ids": [7],
            "bindings": [{"shape_id": 7, "renderer": "image"}],
        })
        missing_title = {
            "top_level_elements": [{
                "shape_id": 100, "type": "group", "children": [{"shape_id": 7}],
                "box": {"x": .1, "y": .2, "width": .8, "height": .6},
            }],
        }
        slots, static = TemplateCatalogBuilder._with_chart_canvas(
            missing_title, [visual], [8, 9], dedicated=True,
        )
        self.assertEqual(slots, [visual])
        self.assertEqual(static, [8, 9])

        low_title = CatalogSlot.model_validate({
            "slot_id": "title", "kind": "text", "role": "title",
            "target_shape_ids": [1],
            "bindings": [{"shape_id": 1, "renderer": "text", "value_paths": ["/text"]}],
        })
        low_slide = {"top_level_elements": [{
            "shape_id": 1, "type": "text", "placeholder_type": "title",
            "box": {"x": .1, "y": .86, "width": .8, "height": .08},
        }]}
        slots, static = TemplateCatalogBuilder._with_chart_canvas(
            low_slide, [low_title], [22], dedicated=True,
        )
        self.assertEqual(slots, [low_title])
        self.assertEqual(static, [22])


class ChartPlannerTests(unittest.TestCase):
    @staticmethod
    def catalog() -> TemplateCatalog:
        title = CatalogSlot.model_validate({
            "slot_id": "title", "kind": "text", "role": "title",
            "target_shape_ids": [10],
            "bindings": [{"shape_id": 10, "renderer": "text", "value_paths": ["/text"]}],
        })
        return TemplateCatalog(
            model="test", source_sha256="hash",
            families=[CatalogFamily(
                family_id="chart", slide_class="chart", description="Chart",
                variants=[CatalogVariant(
                    slide_number=1, description="Line", slots=[title, canvas_slot(preferred=["line"], preserve=[10])],
                    static_shape_ids=[],
                )],
            )],
        )

    @staticmethod
    def asset(rows: int = 12):
        csv = "month,orders,revenue\n" + "\n".join(
            f"{2026 + index // 12}-{index % 12 + 1:02d}-01,{100 + index},{1000 + index * 10}"
            for index in range(rows)
        )
        return parse_csv_bytes(csv.encode(), "Monthly orders")[0].model_copy(
            update={"visual_hint": VisualHint(mode="force", kind="chart", subtype="line")}
        )

    def test_compiler_preserves_clear_targets_during_canvas_cleanup(self):
        asset = self.asset()
        title = self.catalog().families[0].variants[0].slots[0]
        subtitle = CatalogSlot.model_validate({
            "slot_id": "subtitle", "kind": "text", "role": "subtitle",
            "target_shape_ids": [11],
            "bindings": [{"shape_id": 11, "renderer": "text", "value_paths": ["/text"]}],
        })
        variant = CatalogVariant(
            slide_number=1, description="Line",
            slots=[title, subtitle, canvas_slot(preferred=["line"], preserve=[10])],
            static_shape_ids=[],
        )
        slide = SlidePlan.model_validate({
            "number": 1, "purpose": "trend", "template_family_id": "chart",
            "template_slide_number": 1, "template_class": "chart",
            "assignments": [
                {"slot_id": "title", "kind": "text", "target_shape_ids": [10],
                 "action": "replace", "source_refs": ["brief"],
                 "content": {"kind": "text", "text": "Orders"}},
                {"slot_id": "subtitle", "kind": "text", "target_shape_ids": [11],
                 "action": "clear", "source_refs": []},
                {"slot_id": "chart_visual", "kind": "visual", "target_shape_ids": [],
                 "action": "replace", "source_refs": [f"asset:{asset.id}"],
                 "content": {"kind": "visual", "asset_id": asset.id,
                             "visual_kind": "chart", "subtype": "line",
                             "selected_columns": [column.key for column in asset.columns]}},
            ],
        })

        operations = compile_slide(slide, variant, [asset])
        canvas = next(item for item in operations if item.renderer.value == "canvas")

        self.assertEqual(canvas.preserve_shape_ids, [10, 11])

    def test_canonicalization_replaces_model_table_assignment_and_paginates(self):
        asset = self.asset(13)
        request = PlanningRequest(
            brief="Orders", content_package="Monthly trend", structured_assets=[asset],
            slide_count={"min": 1, "max": 2},
        )
        candidate = PlanCandidate.model_validate({
            "deck": {
                "title": {"text": "Orders", "source_refs": ["brief"]},
                "summary": {"text": "Orders", "source_refs": ["brief"]}, "language": "en",
            },
            "slides": [{
                "number": 1, "purpose": "trend", "template_family_id": "chart",
                "template_slide_number": 1, "template_class": "chart",
                "chart_asset_id": asset.id,
                "assignments": [
                    {"slot_id": "title", "kind": "text", "target_shape_ids": [10],
                     "action": "replace", "source_refs": ["brief"],
                     "content": {"kind": "text", "text": "Monthly orders"}},
                    {"slot_id": "chart_visual", "kind": "table", "target_shape_ids": [],
                     "action": "replace", "source_refs": [f"asset:{asset.id}"],
                     "content": {"kind": "table", "columns": ["month"], "rows": [[999]]}},
                ],
            }],
        })
        normalized = ContentPlanner._canonicalize_candidate(self.catalog(), candidate, request)
        assignment = normalized.slides[0].assignments[1]
        self.assertEqual(assignment.kind.value, "visual")
        self.assertEqual(assignment.content.visual_kind, "chart")
        self.assertEqual(assignment.source_refs, [f"asset:{asset.id}"])

        paged = ContentPlanner._paginate_candidate(request, self.catalog(), normalized)
        self.assertEqual([slide.assignments[1].content.page for slide in paged.slides], [1, 2])
        self.assertTrue(paged.slides[1].assignments[0].content.text.endswith("(продолжение)"))
        self.assertEqual(
            [value for slide in paged.slides for value in slide.assignments[1].content.selected_columns],
            [value for _ in paged.slides for value in [column.key for column in asset.columns]],
        )

        analysis = ContentAnalysis.model_validate({
            "goal": "Orders", "audience": "Managers", "language": "en",
            "narrative": ["Orders"], "recommended_slide_count": 2,
            "facts": [{"id": "orders", "text": "Orders", "source_refs": ["brief"]}],
        })
        self.assertEqual(
            validate_candidate(request, [SourceChunk(id="brief", text="Orders"), SourceChunk(
                id=f"asset:{asset.id}", text=json.dumps(asset.model_dump(mode="json")),
            )], analysis, self.catalog(), paged),
            [],
        )

        request.slide_count.max = 1
        with self.assertRaises(PlannerFailure) as error:
            ContentPlanner._paginate_candidate(request, self.catalog(), normalized)
        self.assertEqual(error.exception.code, "content_overflow")

    def test_delivery_service_datasets_become_line_and_donut_slides(self):
        root = Path(__file__).resolve().parents[1] / "description/tasks_examples/delivery_service_results"
        monthly = parse_csv_bytes((root / "monthly_orders.csv").read_bytes(), "Monthly orders")[0]
        regional = parse_csv_bytes((root / "regional_distribution.csv").read_bytes(), "Regional distribution")[0]
        request = PlanningRequest(
            brief="Delivery service results", content_package="Orders and regional mix",
            structured_assets=[monthly, regional], slide_count={"min": 2, "max": 2},
        )

        def slide(number: int, asset_id: str, title: str) -> dict:
            return {
                "number": number, "purpose": title, "template_family_id": "chart",
                "template_slide_number": 1, "template_class": "chart",
                "chart_asset_id": asset_id,
                "assignments": [{
                    "slot_id": "title", "kind": "text", "target_shape_ids": [10],
                    "action": "replace", "source_refs": ["brief"],
                    "content": {"kind": "text", "text": title},
                }],
            }

        candidate = PlanCandidate.model_validate({
            "deck": {
                "title": {"text": "Delivery service", "source_refs": ["brief"]},
                "summary": {"text": "Delivery service", "source_refs": ["brief"]},
                "language": "en",
            },
            "slides": [
                slide(1, monthly.id, "Monthly orders"),
                slide(2, regional.id, "Regional distribution"),
            ],
        })
        normalized = ContentPlanner._canonicalize_candidate(self.catalog(), candidate, request)
        paged = ContentPlanner._paginate_candidate(request, self.catalog(), normalized)
        visuals = [
            next(item.content for item in planned.assignments if item.slot_id == "chart_visual")
            for planned in paged.slides
        ]
        self.assertEqual([(item.subtype, item.page) for item in visuals], [("line", 1), ("donut", 1)])
        self.assertEqual(len(monthly.rows), 12)
        self.assertEqual([row[monthly.columns[0].key] for row in monthly.rows], [
            "January", "February", "March", "April", "May", "June",
            "July", "August", "September", "October", "November", "December",
        ])
        sources = [
            SourceChunk(id="brief", text="Delivery service results"),
            *[
                SourceChunk(id=f"asset:{asset.id}", text=json.dumps(asset.model_dump(mode="json")))
                for asset in (monthly, regional)
            ],
        ]
        analysis = ContentAnalysis.model_validate({
            "goal": "Results", "audience": "Managers", "language": "en",
            "narrative": ["Monthly", "Regional"], "recommended_slide_count": 2,
            "facts": [{"id": "results", "text": "Results", "source_refs": ["brief"]}],
        })
        self.assertEqual(validate_candidate(request, sources, analysis, self.catalog(), paged), [])

    def test_strict_generation_uses_each_chart_once_and_leaves_regular_slide_unassigned(self):
        monthly = parse_csv_bytes(
            b"month,orders\nJan,10\nFeb,12\n", "Monthly orders",
        )[0].model_copy(
            update={"visual_hint": VisualHint(mode="force", kind="chart", subtype="line")}
        )
        regional = parse_csv_bytes(
            b"region,share\nNorth,60\nSouth,40\n", "Regional split",
        )[0].model_copy(
            update={"visual_hint": VisualHint(mode="force", kind="chart", subtype="donut")}
        )
        request = PlanningRequest(
            brief="Delivery results", content_package="Delivery results",
            structured_assets=[monthly, regional], slide_count={"min": 3, "max": 3},
        )
        catalog = self.catalog().model_copy(deep=True)
        title = catalog.families[0].variants[0].slots[0].model_copy(deep=True)
        catalog.families.append(CatalogFamily(
            family_id="content", slide_class="content_text", description="Text",
            variants=[CatalogVariant(
                slide_number=2, description="Text", slots=[title], static_shape_ids=[],
            )],
        ))

        def assignment(text: str) -> list[dict]:
            return [{
                "slot_id": "title", "action": "replace", "source_refs": ["brief"],
                "content": {"kind": "text", "text": text},
            }]

        candidate = {
            "deck": {
                "title": {"text": "Delivery results", "source_refs": ["brief"]},
                "summary": {"text": "Delivery results", "source_refs": ["brief"]},
                "language": "en",
            },
            "slides": [
                {
                    "purpose": "Overview", "template_family_id": "content",
                    "template_slide_number": 2, "chart_asset_id": None,
                    "assignments": assignment("Delivery results"),
                },
                {
                    "purpose": "Monthly trend", "template_family_id": "chart",
                    "template_slide_number": 1, "chart_asset_id": monthly.id,
                    "assignments": assignment("Monthly orders"),
                },
                {
                    "purpose": "Regional split", "template_family_id": "chart",
                    "template_slide_number": 1, "chart_asset_id": regional.id,
                    "assignments": assignment("Regional split"),
                },
            ],
        }
        deck_outline = {
            "deck": candidate["deck"],
            "slides": [{
                "number": index,
                "purpose": slide["purpose"],
                "template_family_id": slide["template_family_id"],
                "template_slide_number": slide["template_slide_number"],
                "source_refs": ["brief", *(
                    [f"asset:{slide['chart_asset_id']}"] if slide["chart_asset_id"] else []
                )],
                "asset_ids": [slide["chart_asset_id"]] if slide["chart_asset_id"] else [],
                "image_ids": [],
                "chart_asset_id": slide["chart_asset_id"],
            } for index, slide in enumerate(candidate["slides"], 1)],
        }
        analysis = {
            "goal": "Delivery results", "audience": "Managers", "language": "en",
            "narrative": ["Overview", "Monthly trend", "Regional split"],
            "recommended_slide_count": 3,
            "facts": [{
                "id": "delivery", "text": "Delivery results",
                "source_refs": ["brief"], "mandatory": True,
            }],
            "visual_opportunities": [], "visual_candidates": [],
        }
        critique = {
            "scores": {
                "factuality": 5, "coverage": 5, "narrative": 5,
                "template_fit": 5, "conciseness": 5,
            },
            "issues": [],
        }

        class NativeLLM:
            model = "native-test"

            def complete(self, prompt, **kwargs):
                if kwargs["stage"] == "content analysis":
                    return json.dumps(analysis)
                if kwargs["stage"] == "deck outline":
                    return json.dumps(deck_outline)
                if kwargs["stage"] == "slide draft":
                    match = re.search(r'SLIDE OUTLINE:\n(\{[^\n]+\})', prompt)
                    number = json.loads(match.group(1))["number"]
                    return json.dumps({
                        "number": number,
                        "assignments": candidate["slides"][number - 1]["assignments"],
                    })
                if kwargs["stage"] == "plan critique":
                    return json.dumps(critique)
                raise AssertionError(kwargs["stage"])

        plan = ContentPlanner(
            llm=NativeLLM(), catalog=catalog, max_revisions=0,
        ).generate(request)

        self.assertEqual(plan.quality.status, "passed")
        self.assertEqual(plan.slides[0].chart_asset_id, None)
        self.assertEqual(
            [slide.chart_asset_id for slide in plan.slides if slide.chart_asset_id],
            [monthly.id, regional.id],
        )


class ChartCanvasBuilderTests(unittest.TestCase):
    def test_cleans_body_and_adds_exactly_one_native_chart(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "template"
            template.mkdir()
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title_group = slide.shapes.add_group_shape()
            title = title_group.shapes.add_textbox(Inches(.5), Inches(.3), Inches(8), Inches(.7))
            title.text = "Demo title"
            body = slide.shapes.add_textbox(Inches(.5), Inches(1.7), Inches(5), Inches(2))
            body.text = "Old demo body"
            logo = slide.shapes.add_textbox(Inches(8.8), Inches(.2), Inches(.8), Inches(.4))
            logo.text = "BRAND"
            logo.name = "Brand logo"
            footer = slide.shapes.add_textbox(Inches(.4), Inches(7.1), Inches(2), Inches(.25))
            footer.text = "Footer"
            source = template / "source.pptx"
            presentation.save(source)
            source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            (template / "presentation.json").write_text(json.dumps({"source_sha256": source_hash}))

            title_slot = CatalogSlot.model_validate({
                "slot_id": "title", "kind": "text", "role": "title",
                "target_shape_ids": [title.shape_id],
                "bindings": [{"shape_id": title.shape_id, "renderer": "text", "value_paths": ["/text"]}],
            })
            variant = CatalogVariant(
                slide_number=1, description="Chart",
                slots=[title_slot, canvas_slot(preferred=["line"], preserve=[title.shape_id, logo.shape_id, footer.shape_id])],
                static_shape_ids=[logo.shape_id, footer.shape_id],
            )
            catalog = TemplateCatalog(
                model="test", source_sha256=source_hash,
                families=[CatalogFamily(
                    family_id="chart", slide_class="chart", description="Chart", variants=[variant],
                )],
            )
            (template / "planner_catalog.json").write_text(catalog.model_dump_json(indent=2))
            asset = ChartPlannerTests.asset()
            slide_plan = SlidePlan.model_validate({
                "number": 1, "purpose": "trend", "template_family_id": "chart",
                "template_slide_number": 1, "template_class": "chart", "chart_asset_id": asset.id,
                "assignments": [
                    {"slot_id": "title", "kind": "text", "target_shape_ids": [title.shape_id],
                     "action": "replace", "source_refs": ["brief"],
                     "content": {"kind": "text", "text": "Orders trend"}},
                    {"slot_id": "chart_visual", "kind": "visual", "target_shape_ids": [],
                     "action": "replace", "source_refs": [f"asset:{asset.id}"],
                     "content": {"kind": "visual", "asset_id": asset.id, "visual_kind": "chart",
                                 "subtype": "line", "selected_columns": [column.key for column in asset.columns]}},
                ],
            })
            slide_plan.render_operations = compile_slide(slide_plan, variant, [asset])
            self.assertEqual(slide_plan.render_operations[-1].renderer.value, "canvas")
            self.assertEqual(slide_plan.render_operations[-1].asset_hash, hashlib.sha256(
                json.dumps(asset.model_dump(mode="json"), ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
            ).hexdigest())
            plan = PresentationPlan(
                model="test", template=TemplateReference(
                    source_sha256=source_hash, catalog_version=CATALOG_SCHEMA_VERSION,
                ),
                deck=DeckInfo(
                    title=SourcedText(text="Orders", source_refs=["brief"]),
                    summary=SourcedText(text="Orders", source_refs=["brief"]), language="en",
                ),
                sources=[SourceChunk(id="brief", text="Orders")], structured_assets=[asset],
                slides=[slide_plan], quality=QualityReport(status="passed", revision_count=0),
            )
            output = root / "result.pptx"
            PresentationBuilder(libreoffice_path="missing").build(template, plan, output)
            result = Presentation(output).slides[0]
            charts = [shape for shape in result.shapes if shape.has_chart]
            self.assertEqual(len(charts), 1)
            self.assertEqual([category.label for category in charts[0].chart.plots[0].categories], [
                row[asset.columns[0].key] for row in asset.rows
            ])
            self.assertEqual([series.name for series in charts[0].chart.series], ["orders", "revenue"])
            self.assertFalse(charts[0].chart.has_title)
            self.assertTrue(charts[0].chart.has_legend)
            def texts_in(shapes):
                values = []
                for shape in shapes:
                    if getattr(shape, "has_text_frame", False):
                        values.append(shape.text)
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        values.extend(texts_in(shape.shapes))
                return values

            texts = texts_in(result.shapes)
            self.assertIn("Orders trend", texts)
            self.assertIn("BRAND", texts)
            self.assertIn("Footer", texts)
            self.assertNotIn("Old demo body", texts)


if __name__ == "__main__":
    unittest.main()
