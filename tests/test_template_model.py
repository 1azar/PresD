from __future__ import annotations

import hashlib
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from template_model import SlideClass, TemplateModel, VLMStatus
from template_model.analyzer import extract_elements, render_pdf
from template_model.main import main as cli_main
from template_model.models import Box, SlideElement
from template_model.structures import detect_structures
from template_model.vlm import VLMAnalyzer


class FakeResponses:
    def __init__(self, outputs):
        self.outputs = iter(outputs)
        self.calls = 0

    def create(self, **kwargs):
        self.calls += 1
        value = next(self.outputs)
        if isinstance(value, Exception):
            raise value
        return SimpleNamespace(output_text=value)


class FakeClient:
    def __init__(self, outputs):
        self.responses = FakeResponses(outputs)


class FakeChatCompletions:
    def __init__(self, outputs):
        self.outputs = iter(outputs)
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        value = next(self.outputs)
        if isinstance(value, Exception):
            raise value
        return {"choices": [{"finish_reason": "stop", "message": {"content": value}}]}


class FakeChatClient:
    def __init__(self, outputs):
        self.completions = FakeChatCompletions(outputs)
        self.chat = SimpleNamespace(completions=self.completions)


def fake_render(source: Path, destination: Path, libreoffice: str, pdftoppm: str) -> list[Path]:
    slide_count = len(Presentation(source).slides)
    destination.mkdir(parents=True, exist_ok=True)
    result = []
    for index in range(1, slide_count + 1):
        path = destination / f"slide_{index:03d}.png"
        Image.new("RGB", (32, 18), "white").save(path)
        result.append(path)
    return result


class PdfRendererTests(unittest.TestCase):
    def test_renders_valid_pdf_atomically_with_isolated_profile(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.pptx"
            destination = root / "presentation.pdf"
            source.write_bytes(b"pptx")
            commands: list[list[str]] = []

            def convert(command, **_kwargs):
                commands.append(command)
                output = Path(command[command.index("--outdir") + 1]) / "source.pdf"
                output.write_bytes(b"%PDF-1.7\nresult")

            with patch("template_model.analyzer.subprocess.run", side_effect=convert):
                result = render_pdf(source, destination, "libreoffice-test")

            self.assertEqual(result, destination.resolve())
            self.assertEqual(destination.read_bytes(), b"%PDF-1.7\nresult")
            self.assertEqual(commands[0][0], "libreoffice-test")
            self.assertTrue(any(item.startswith("-env:UserInstallation=file:") for item in commands[0]))

    def test_invalid_conversion_does_not_replace_existing_pdf(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.pptx"
            destination = root / "presentation.pdf"
            source.write_bytes(b"pptx")
            destination.write_bytes(b"%PDF-existing")

            def convert(command, **_kwargs):
                output = Path(command[command.index("--outdir") + 1]) / "source.pdf"
                output.write_bytes(b"invalid")

            with patch("template_model.analyzer.subprocess.run", side_effect=convert):
                with self.assertRaisesRegex(RuntimeError, "invalid PDF"):
                    render_pdf(source, destination)

            self.assertEqual(destination.read_bytes(), b"%PDF-existing")


class ParserTests(unittest.TestCase):
    @staticmethod
    def text_element(shape_id: int, x: float, y: float, text: str = "Text") -> SlideElement:
        return SlideElement(
            shape_id=shape_id, name=str(shape_id), type="text",
            box=Box(x=x, y=y, width=0.18, height=0.04), z_order=shape_id, text=text,
        )

    def test_detects_pseudo_table_as_ordered_grid(self):
        elements = [
            self.text_element(row * 3 + column + 1, 0.1 + column * 0.25, 0.2 + row * 0.08)
            for row in range(4) for column in range(3)
        ]
        structure = detect_structures(elements)[0]
        self.assertEqual(structure.kind, "grid")
        self.assertEqual((structure.rows, structure.columns, structure.header_rows), (4, 3, 1))
        self.assertEqual([cell.shape_ids[0] for cell in structure.cells], list(range(1, 13)))

    def test_detects_repeat_and_rejects_irregular_layout(self):
        elements = [
            self.text_element(item * 2 + field + 1, 0.1 + item * 0.28, 0.3 + field * 0.12)
            for item in range(3) for field in range(2)
        ]
        repeat = detect_structures(elements)[0]
        self.assertEqual(repeat.kind, "repeat")
        self.assertEqual(len(repeat.items), 3)
        self.assertEqual([field.field_id for field in repeat.fields], ["field_0", "field_1"])
        elements[-1].box.x = 0.55
        self.assertEqual(detect_structures(elements), [])

    def test_detects_composite_static_visual(self):
        image = SlideElement(
            shape_id=1, name="diagram", type="image",
            box=Box(x=0.2, y=0.25, width=0.6, height=0.6), z_order=0,
        )
        label = self.text_element(2, 0.35, 0.4, "Annotation")
        structures = detect_structures([image, label], "content_visual")
        self.assertEqual(structures[0].kind, "static_visual")
        self.assertEqual(structures[0].shape_ids, [1, 2])

    def test_extracts_content_and_omits_decorative_shapes(self):
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "image.png"
            Image.new("RGB", (200, 100), "red").save(image_path)
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            text = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(4), Inches(1))
            text.text = "Exact\nslide text"
            slide.shapes.add_picture(str(image_path), Inches(1), Inches(2), Inches(2), Inches(1))
            table = slide.shapes.add_table(2, 2, Inches(4), Inches(2), Inches(4), Inches(2))
            table.table.cell(0, 0).text = "Header"
            table.table.cell(1, 1).text = "Value"
            data = ChartData()
            data.categories = ["A", "B"]
            data.add_series("Revenue", (1, 2))
            slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED,
                Inches(1), Inches(4), Inches(5), Inches(2), data,
            )
            slide.shapes.add_shape(1, Inches(8), Inches(4), Inches(1), Inches(1))

            elements = extract_elements(slide, prs.slide_width, prs.slide_height)
            kinds = [element.type for element in elements]
            self.assertEqual(kinds.count("text"), 1)
            self.assertTrue({"image", "table", "chart"} <= set(kinds))
            self.assertEqual(len(elements), 4)
            self.assertEqual(next(e for e in elements if e.type == "text").text, "Exact\nslide text")
            table_element = next(e for e in elements if e.type == "table")
            self.assertEqual((table_element.table.rows, table_element.table.columns), (2, 2))
            self.assertEqual(table_element.table.cells[0].text, "Header")
            chart = next(e for e in elements if e.type == "chart").chart
            self.assertEqual(chart.series_names, ["Revenue"])
            self.assertEqual(chart.category_labels, ["A", "B"])
            self.assertTrue(all(
                0 <= coordinate <= 1
                for element in elements
                for coordinate in element.box.model_dump().values()
            ))

    def test_group_hierarchy_and_placeholder_type(self):
        child = SimpleNamespace(
            shape_id=2,
            name="child",
            shape_type=MSO_SHAPE_TYPE.TEXT_BOX,
            has_table=False,
            has_chart=False,
            has_text_frame=True,
            text="inside",
            is_placeholder=False,
            left=10,
            top=10,
            width=20,
            height=20,
        )
        group = SimpleNamespace(
            shape_id=1,
            name="group",
            shape_type=MSO_SHAPE_TYPE.GROUP,
            has_table=False,
            has_chart=False,
            has_text_frame=False,
            is_placeholder=False,
            left=0,
            top=0,
            width=50,
            height=50,
            shapes=[child],
        )
        slide = SimpleNamespace(shapes=[group])
        elements = extract_elements(slide, 100, 100)
        self.assertEqual(elements[0].type, "group")
        self.assertEqual(elements[0].children[0].shape_id, 2)
        self.assertEqual(elements[0].children[0].text, "inside")


class VLMTests(unittest.TestCase):
    valid = json.dumps({
        "classification": "cover",
        "tags": ["Two Columns", "two-columns", "Центр"],
        "description": "Титульный слайд.",
        "confidence": 0.9,
    })

    def run_analyzer(self, outputs):
        temporary = tempfile.TemporaryDirectory()
        image = Path(temporary.name) / "slide.png"
        Image.new("RGB", (10, 10), "white").save(image)
        client = FakeClient(outputs)
        return temporary, image, client, VLMAnalyzer(client=client)

    def test_valid_result_and_tag_normalization(self):
        temporary, image, client, analyzer = self.run_analyzer([self.valid])
        with temporary:
            result = analyzer.analyze(image, {"elements": []})
        self.assertEqual(result.classification, SlideClass.COVER)
        self.assertEqual(result.tags, ["two_columns", "центр"])
        self.assertEqual(client.responses.calls, 1)

    def test_invalid_json_and_timeout_retry(self):
        for outputs, error in [
            (["not-json"] * 3, "Invalid JSON"),
            ([TimeoutError("late")] * 3, "late"),
        ]:
            temporary, image, client, analyzer = self.run_analyzer(outputs)
            with temporary, self.assertRaisesRegex(Exception, error):
                analyzer.analyze(image, {"elements": []})
            self.assertEqual(client.responses.calls, 3)

    def test_chat_completions_uses_multimodal_message(self):
        temporary = tempfile.TemporaryDirectory()
        image = Path(temporary.name) / "slide.png"
        Image.new("RGB", (10, 10), "white").save(image)
        client = FakeChatClient([self.valid])
        analyzer = VLMAnalyzer(client=client, model="vision-model")

        with temporary:
            result = analyzer.analyze(image, {"elements": []})

        self.assertEqual(result.classification, SlideClass.COVER)
        content = client.completions.calls[0]["messages"][0]["content"]
        self.assertEqual(content[0]["type"], "text")
        self.assertEqual(content[1]["type"], "image_url")
        self.assertTrue(content[1]["image_url"]["url"].startswith("data:image/png;base64,"))


class EndToEndTests(unittest.TestCase):
    @staticmethod
    def make_source(path: Path, slide_count: int = 1) -> None:
        prs = Presentation()
        for index in range(slide_count):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = f"Title {index + 1}"
            slide.placeholders[1].text = f"Body {index + 1}"
        prs.save(path)

    def test_analyze_creates_per_slide_contract_without_vlm(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.pptx"
            output = root / "result"
            self.make_source(source, 2)
            before = hashlib.sha256(source.read_bytes()).hexdigest()

            with patch("template_model.analyzer.render_previews", side_effect=fake_render):
                manifest = TemplateModel().analyze(source, output)

            saved = json.loads((output / "presentation.json").read_text(encoding="utf-8"))
            self.assertEqual(saved["schema_version"], "3.0")
            self.assertEqual(saved["slide_count"], 2)
            self.assertFalse(saved["vlm_enabled"])
            self.assertEqual(manifest.slides[0].vlm_status, VLMStatus.DISABLED)
            for index in (1, 2):
                slide_dir = output / "slides" / f"slide_{index:03d}"
                metadata = json.loads((slide_dir / "metadata.json").read_text(encoding="utf-8"))
                self.assertTrue((slide_dir / "preview.png").is_file())
                self.assertEqual(metadata["vlm"]["status"], "disabled")
                self.assertEqual([e["text"] for e in metadata["elements"]], [f"Title {index}", f"Body {index}"])
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
            self.assertEqual(hashlib.sha256((output / "source.pptx").read_bytes()).hexdigest(), before)
            self.assertFalse((output / ".rendered_previews").exists())

    def test_renderer_and_vlm_failures_are_nonfatal(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.pptx"
            output = root / "result"
            self.make_source(source)
            manifest = TemplateModel(
                use_vlm=True,
                vlm_client=FakeClient([]),
                libreoffice_path="missing-libreoffice",
            ).analyze(source, output)
            self.assertEqual(manifest.slides[0].vlm_status, VLMStatus.ERROR)
            self.assertIsNone(manifest.slides[0].preview_path)
            metadata = json.loads(
                (output / "slides" / "slide_001" / "metadata.json").read_text(encoding="utf-8")
            )
            self.assertIn("Preview is unavailable", metadata["vlm"]["error"])
            self.assertTrue(manifest.warnings)

    def test_vlm_success_is_saved(self):
        response = json.dumps({
            "classification": "content_text",
            "tags": ["title", "two columns"],
            "description": "Текстовый слайд.",
            "confidence": 0.8,
        })
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.pptx"
            output = root / "result"
            self.make_source(source)
            with patch("template_model.analyzer.render_previews", side_effect=fake_render):
                manifest = TemplateModel(use_vlm=True, vlm_client=FakeClient([response])).analyze(
                    source, output
                )
            self.assertEqual(manifest.slides[0].vlm_status, VLMStatus.OK)
            metadata = json.loads(
                (output / "slides" / "slide_001" / "metadata.json").read_text(encoding="utf-8")
            )
            self.assertEqual(metadata["vlm"]["classification"], "content_text")
            self.assertEqual(metadata["vlm"]["tags"], ["title", "two_columns"])

    def test_vlm_failure_does_not_stop_following_slides(self):
        response = json.dumps({
            "classification": "closing",
            "tags": ["contacts"],
            "description": "Финальный слайд.",
            "confidence": 0.9,
        })
        client = FakeClient(["invalid", "invalid", "invalid", response])
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.pptx"
            output = root / "result"
            self.make_source(source, 2)
            with patch("template_model.analyzer.render_previews", side_effect=fake_render):
                manifest = TemplateModel(use_vlm=True, vlm_client=client).analyze(source, output)
            self.assertEqual(
                [entry.vlm_status for entry in manifest.slides],
                [VLMStatus.ERROR, VLMStatus.OK],
            )
            self.assertEqual(client.responses.calls, 4)

    def test_cli_uses_structural_mode_by_default(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.pptx"
            output = root / "result"
            self.make_source(source)
            stdout = io.StringIO()
            with patch("template_model.analyzer.render_previews", side_effect=fake_render), redirect_stdout(stdout):
                code = cli_main([
                    "analyze",
                    "--template", str(source),
                    "--output", str(output),
                ])
            summary = json.loads(stdout.getvalue())
            self.assertEqual(code, 0)
            self.assertFalse(summary["vlm_enabled"])
            self.assertEqual(summary["slides"], 1)

    def test_nonempty_output_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.pptx"
            output = root / "result"
            self.make_source(source)
            output.mkdir()
            (output / "keep.txt").write_text("do not overwrite", encoding="utf-8")
            with self.assertRaises(FileExistsError):
                TemplateModel().analyze(source, output)
            self.assertEqual((output / "keep.txt").read_text(encoding="utf-8"), "do not overwrite")


if __name__ == "__main__":
    unittest.main()
