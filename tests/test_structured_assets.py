from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from content_planner.models import (
    CatalogBinding,
    CatalogSlot,
    DatasetAsset,
    DiagramAsset,
    PictogramGridAsset,
    SlotAssignment,
    VisualCapabilities,
    VisualHint,
)
from content_planner.compiler import CompilationError, compile_assignment
from content_planner.structured_assets import (
    StructuredAssetError,
    asset_hash,
    choose_visual,
    materialize_visual,
    paginate_asset,
    parse_csv_bytes,
    parse_docx_tables,
    parse_json_assets,
    parse_markdown_tables,
    parse_xlsx,
)


class StructuredImportTests(unittest.TestCase):
    def test_csv_bom_delimiter_types_and_stable_id(self):
        raw = "\ufeffМесяц;Выручка;Активен\nЯнварь;120;true\nФевраль;145;false\n".encode()
        first = parse_csv_bytes(raw, "Выручка")[0]
        second = parse_csv_bytes(raw, "Выручка")[0]
        self.assertEqual(first.id, second.id)
        self.assertEqual([column.type for column in first.columns], ["text", "number", "boolean"])
        self.assertEqual(first.rows[1]["выручка"], 145)

    def test_xlsx_uses_visible_nonempty_sheets(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Visible"
        sheet.append(["Date", "Value"])
        sheet.append(["2026-01-01", 4])
        hidden = workbook.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        hidden.append(["A", "B"])
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        assets = parse_xlsx(stream)
        self.assertEqual([asset.title for asset in assets], ["Visible"])
        self.assertEqual(assets[0].columns[0].type, "date")

    def test_markdown_and_docx_tables_are_datasets(self):
        markdown = parse_markdown_tables("# X\n\n| Name | Value |\n|---|---:|\n| A | 2 |")
        self.assertEqual(markdown[0].rows[0]["value"], 2)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.docx"
            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Value"
            table.cell(1, 0).text, table.cell(1, 1).text = "A", "3"
            document.save(path)
            assets = parse_docx_tables(path)
        self.assertEqual(assets[0].rows[0]["value"], 3)

    def test_json_is_strictly_typed(self):
        valid = '{"id":"d","kind":"dataset","title":"D","columns":[{"key":"x","label":"X","type":"number"}],"rows":[{"x":1}]}'
        self.assertIsInstance(parse_json_assets(valid)[0], DatasetAsset)
        with self.assertRaises(StructuredAssetError):
            parse_json_assets(valid[:-1] + ',"unexpected":1}')


class VisualSelectionTests(unittest.TestCase):
    def dataset(self, count: int = 7) -> DatasetAsset:
        return parse_csv_bytes(
            ("Category,Value\n" + "\n".join(f"Item {index},{index}" for index in range(count))).encode(),
            "Values",
        )[0]

    def test_auto_force_and_none(self):
        asset = self.dataset(7)
        self.assertEqual(choose_visual(asset), ("chart", "column"))
        forced = asset.model_copy(update={"visual_hint": VisualHint(mode="force", kind="table")})
        self.assertEqual(choose_visual(forced), ("table", None))
        disabled = asset.model_copy(update={"visual_hint": VisualHint(mode="none")})
        self.assertIsNone(choose_visual(disabled))

    def test_time_series_scatter_and_no_implicit_sort(self):
        dates = parse_csv_bytes(b"Date,Value\n2026-02-01,2\n2026-01-01,1", "Time")[0]
        self.assertEqual(choose_visual(dates), ("chart", "line"))
        selection = paginate_asset(dates, "chart", "line", None)[0]
        rendered = materialize_visual(dates, selection)
        self.assertEqual(rendered["categories"], ["2026-02-01", "2026-01-01"])
        scatter = parse_csv_bytes(b"X,Y\n1,4\n2,3", "XY")[0]
        self.assertEqual(choose_visual(scatter), ("chart", "scatter"))

    def test_month_percentage_and_capability_fallback(self):
        monthly = parse_csv_bytes(
            b"month,orders_2024,orders_2025\nJanuary,2,3\nFebruary,4,5", "Orders"
        )[0]
        self.assertEqual(choose_visual(monthly), ("chart", "line"))
        regions = parse_csv_bytes(
            b"region,orders,share_percent\nMoscow,80,40\nOther,120,60", "Regions"
        )[0]
        self.assertEqual(choose_visual(regions), ("chart", "donut"))
        table_only = VisualCapabilities(
            kinds=["table"], render_mode="container", width=1, height=1,
            aspect_ratio=1, max_rows=9, max_columns=3,
        )
        self.assertEqual(choose_visual(monthly, table_only), ("table", None))
        forced = monthly.model_copy(update={
            "visual_hint": VisualHint(mode="force", kind="chart", subtype="line")
        })
        with self.assertRaisesRegex(StructuredAssetError, "incompatible"):
            choose_visual(forced, table_only)

    def test_table_pagination_repeats_identifier_and_loses_nothing(self):
        asset = parse_csv_bytes(
            b"ID,A,B,C\n1,10,11,12\n2,20,21,22\n3,30,31,32", "Wide"
        )[0]
        caps = VisualCapabilities(
            kinds=["table"], render_mode="container", width=1, height=1,
            aspect_ratio=1, max_rows=2, max_columns=3,
        )
        pages = paginate_asset(asset, "table", None, caps)
        self.assertEqual(len(pages), 4)
        self.assertTrue(all(page.selected_columns[0] == "id" for page in pages))
        values = [materialize_visual(asset, page, caps)["rows"] for page in pages]
        self.assertEqual(sum(len(page) for page in values), 6)
        with self.assertRaisesRegex(StructuredAssetError, "content_overflow"):
            paginate_asset(asset, "table", None, caps, max_slides=3)

    def test_diagram_and_pictogram_contracts(self):
        diagram = DiagramAsset.model_validate({
            "id": "flow", "kind": "diagram", "title": "Flow", "diagram_type": "process",
            "nodes": [{"id": "a", "label": "A"}, {"id": "b", "label": "B"}],
            "edges": [{"source": "a", "target": "b"}], "order": ["a", "b"],
        })
        self.assertEqual(choose_visual(diagram), ("diagram", "process"))
        pictograms = PictogramGridAsset.model_validate({
            "id": "icons", "kind": "pictogram_grid", "title": "Icons",
            "items": [{"label": "Date", "value": 2, "icon_query": "calendar"}],
        })
        self.assertEqual(choose_visual(pictograms), ("pictogram_grid", None))


class VisualCompilerTests(unittest.TestCase):
    def test_compiler_materializes_asset_and_hash(self):
        asset = parse_csv_bytes(b"Category,Value\nA,1\nB,2", "Data")[0]
        asset = asset.model_copy(update={"visual_hint": VisualHint(mode="force", kind="chart", subtype="column")})
        slot = CatalogSlot(
            slot_id="visual", kind="visual", role="visual", target_shape_ids=[7],
            bindings=[CatalogBinding(shape_id=7, renderer="visual")],
            visual_capabilities=VisualCapabilities(
                kinds=["chart", "table"], render_mode="container", width=1, height=1,
                aspect_ratio=1, max_items=10,
            ),
        )
        assignment = SlotAssignment.model_validate({
            "slot_id": "visual", "kind": "visual", "target_shape_ids": [7],
            "action": "replace", "source_refs": [f"asset:{asset.id}"],
            "content": {"kind": "visual", "asset_id": asset.id, "visual_kind": "chart", "subtype": "column"},
        })
        operation = compile_assignment(slot, assignment, {asset.id: asset})[0]
        self.assertEqual(operation.asset_hash, asset_hash(asset))
        self.assertEqual(operation.data["series"][0]["values"], [1, 2])
        with self.assertRaises(CompilationError):
            compile_assignment(slot, assignment, {})


if __name__ == "__main__":
    unittest.main()
