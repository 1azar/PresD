from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from content_planner.compiler import compile_assignment
from content_planner.models import (
    CATALOG_SCHEMA_VERSION,
    CatalogBinding,
    CatalogFamily,
    CatalogSlot,
    CatalogVariant,
    DeckInfo,
    PresentationPlan,
    QualityReport,
    RenderOperation,
    SlidePlan,
    SlotAssignment,
    SourceChunk,
    SourcedText,
    TemplateReference,
    TemplateCatalog,
)
from content_planner.compiler import compile_slide
from presentation_builder import BuildFailure, PresentationBuilder
from presentation_builder.core import _create_chart


class CompilerTests(unittest.TestCase):
    def test_compiles_structured_content_and_clears_missing_items(self):
        slot = CatalogSlot.model_validate({
            "slot_id": "metrics",
            "kind": "metrics",
            "role": "metrics",
            "target_shape_ids": [10, 11, 12],
            "capacity": 2,
            "bindings": [
                {"shape_id": 10, "renderer": "text", "value_paths": ["/items/0/value", "/items/0/unit"], "separator": " "},
                {"shape_id": 11, "renderer": "text", "value_paths": ["/items/0/label"]},
                {"shape_id": 12, "renderer": "text", "value_paths": ["/items/1/label"]},
            ],
        })
        assignment = SlotAssignment.model_validate({
            "slot_id": "metrics",
            "kind": "metrics",
            "target_shape_ids": [10, 11, 12],
            "action": "replace",
            "source_refs": ["content_001"],
            "content": {"kind": "metrics", "items": [{"label": "MTTR", "value": 34, "unit": "мин"}]},
        })
        operations = compile_assignment(slot, assignment)
        self.assertEqual(operations[0].text, ["34 мин"])
        self.assertEqual(operations[1].text, ["MTTR"])
        self.assertEqual(operations[2].action.value, "clear")


class CLITests(unittest.TestCase):
    def test_main_file_can_be_executed_directly(self):
        project_root = Path(__file__).resolve().parents[1]
        result = subprocess.run(
            [sys.executable, str(project_root / "presentation_builder" / "main.py"), "--help"],
            cwd=project_root,
            capture_output=True,
            text=True,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Assemble a PPTX", result.stdout)


class BuilderTests(unittest.TestCase):
    def test_image_anchor_is_replaced_by_native_chart_with_same_geometry(self):
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "anchor.png"
            Image.new("RGB", (20, 20), "white").save(image_path)
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            anchor = slide.shapes.add_picture(
                str(image_path), Inches(1.2), Inches(2.1), Inches(5.4), Inches(2.7)
            )
            geometry = (anchor.left, anchor.top, anchor.width, anchor.height)
            chart = _create_chart(slide, anchor, {
                "chart_type": "line", "categories": ["Jan", "Feb"],
                "series": [{"name": "Orders", "values": [10, 12]}],
            })
            self.assertTrue(chart.has_chart)
            self.assertEqual((chart.left, chart.top, chart.width, chart.height), geometry)
            self.assertEqual(len(chart.chart.series), 1)

    def create_template(self, root: Path) -> tuple[Path, str, list[int]]:
        template = root / "template"
        template.mkdir()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Old title"
        slide.placeholders[1].text = "Old body"
        image_path = root / "old.png"
        Image.new("RGB", (40, 20), "red").save(image_path)
        picture = slide.shapes.add_picture(str(image_path), Inches(1), Inches(3), Inches(2), Inches(1))
        shape_ids = [slide.shapes.title.shape_id, slide.placeholders[1].shape_id, picture.shape_id]
        source = template / "source.pptx"
        presentation.save(source)
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        (template / "presentation.json").write_text(
            json.dumps({"source_sha256": source_hash}), encoding="utf-8"
        )
        catalog = TemplateCatalog(
            model="test",
            source_sha256=source_hash,
            families=[CatalogFamily.model_validate({
                "family_id": "cover",
                "slide_class": "cover",
                "description": "Test cover",
                "variants": [{
                    "slide_number": 1,
                    "description": "Test slide",
                    "slots": [
                        {"slot_id": "title", "kind": "text", "role": "title", "target_shape_ids": [shape_ids[0]], "bindings": [{"shape_id": shape_ids[0], "renderer": "text", "value_paths": ["/text"]}]},
                        {"slot_id": "body", "kind": "bullet_list", "role": "body", "target_shape_ids": [shape_ids[1]], "bindings": [{"shape_id": shape_ids[1], "renderer": "text", "value_paths": ["/items"], "mode": "paragraphs"}]},
                        {"slot_id": "picture", "kind": "image", "role": "picture", "target_shape_ids": [shape_ids[2]], "required": False, "bindings": [{"shape_id": shape_ids[2], "renderer": "image"}]},
                    ],
                    "static_shape_ids": [],
                }],
            })],
        )
        (template / "planner_catalog.json").write_text(
            catalog.model_dump_json(indent=2), encoding="utf-8"
        )
        return template, source_hash, shape_ids

    @staticmethod
    def plan(source_hash: str, shape_ids: list[int], asset: str | None = None) -> PresentationPlan:
        variant = CatalogVariant.model_validate({
            "slide_number": 1,
            "description": "Test slide",
            "slots": [
                {"slot_id": "title", "kind": "text", "role": "title", "target_shape_ids": [shape_ids[0]], "bindings": [{"shape_id": shape_ids[0], "renderer": "text", "value_paths": ["/text"]}]},
                {"slot_id": "body", "kind": "bullet_list", "role": "body", "target_shape_ids": [shape_ids[1]], "bindings": [{"shape_id": shape_ids[1], "renderer": "text", "value_paths": ["/items"], "mode": "paragraphs"}]},
                {"slot_id": "picture", "kind": "image", "role": "picture", "target_shape_ids": [shape_ids[2]], "required": False, "bindings": [{"shape_id": shape_ids[2], "renderer": "image"}]},
            ],
            "static_shape_ids": [],
        })
        slides = []
        for number, title in ((1, "First"), (2, "Second")):
            assignments = [
                SlotAssignment.model_validate({"slot_id": "title", "kind": "text", "target_shape_ids": [shape_ids[0]], "action": "replace", "source_refs": ["brief"], "content": {"kind": "text", "text": title}}),
                SlotAssignment.model_validate({"slot_id": "body", "kind": "bullet_list", "target_shape_ids": [shape_ids[1]], "action": "replace", "source_refs": ["brief"], "content": {"kind": "bullet_list", "items": ["A", "B"]}}),
            ]
            if asset:
                assignments.append(SlotAssignment.model_validate({"slot_id": "picture", "kind": "image", "target_shape_ids": [shape_ids[2]], "action": "replace", "source_refs": ["brief"], "content": {"kind": "image", "asset_ref": asset}}))
            else:
                assignments.append(SlotAssignment.model_validate({"slot_id": "picture", "kind": "image", "target_shape_ids": [shape_ids[2]], "action": "keep", "source_refs": [], "content": None}))
            slide = SlidePlan(
                number=number,
                purpose="test",
                template_family_id="cover",
                template_slide_number=1,
                template_class="cover",
                assignments=assignments,
            )
            slide.render_operations = compile_slide(slide, variant)
            slides.append(slide)
        return PresentationPlan(
            model="test",
            template=TemplateReference(
                source_sha256=source_hash, catalog_version=CATALOG_SCHEMA_VERSION
            ),
            deck=DeckInfo(
                title=SourcedText(text="Test", source_refs=["brief"]),
                summary=SourcedText(text="Test", source_refs=["brief"]),
                language="en",
            ),
            sources=[SourceChunk(id="brief", text=f"Test {asset or ''}".strip())],
            slides=slides,
            quality=QualityReport(status="passed", revision_count=0),
        )

    def test_clones_repeated_slide_and_applies_independent_edits(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template, source_hash, ids = self.create_template(root)
            replacement = root / "assets" / "new.png"
            replacement.parent.mkdir()
            Image.new("RGB", (20, 40), "blue").save(replacement)
            output = root / "result.pptx"
            report = PresentationBuilder(libreoffice_path="missing").build(
                template, self.plan(source_hash, ids, "new.png"), output,
                assets_dir=replacement.parent,
            )
            result = Presentation(output)
            self.assertEqual(len(result.slides), 2)
            self.assertEqual([slide.shapes.title.text for slide in result.slides], ["First", "Second"])
            self.assertEqual(result.slides[0].placeholders[1].text, "A\nB")
            self.assertTrue((root / "result.qa" / "report.json").is_file())
            self.assertTrue(any(issue.code == "preview_failed" for issue in report.issues))

    def test_missing_asset_is_atomic(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template, source_hash, ids = self.create_template(root)
            output = root / "result.pptx"
            with self.assertRaises(BuildFailure) as context:
                PresentationBuilder(libreoffice_path="missing").build(
                    template, self.plan(source_hash, ids, "missing.png"), output,
                    assets_dir=root,
                )
            self.assertEqual(context.exception.code, "render_operation_failed")
            self.assertFalse(output.exists())
            self.assertFalse((root / "result.qa").exists())


if __name__ == "__main__":
    unittest.main()
